import io

import pandas as pd
from openpyxl import load_workbook
from PIL import Image

from core.reports.exporters import to_excel, to_pdf
from core.services import settings as svc


def _logo() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (120, 60), (31, 78, 121)).save(buf, format="PNG")
    return buf.getvalue()


def test_pdf_and_excel_with_logo_footer_and_totals(db, master):
    svc.update_settings(master, {"footer_text": "سطر الفوتر الأول\nسطر تاني", "signature_labels": "المحاسب,المدير"})
    svc.set_logo(master, _logo(), "image/png")
    settings = svc.get_settings()
    df = pd.DataFrame({"البند": ["خامات", "تشغيل طويل جداً " * 6], "القيمة": [4.6, 2.5], "نشط": [True, False]})

    pdf = to_pdf(df, "تقرير تجربة", settings, subtitle="فترة", printed_by="Admin", totals={"القيمة": 7.1})
    assert pdf.startswith(b"%PDF")

    wb = load_workbook(io.BytesIO(to_excel(df, "تقرير تجربة", settings, printed_by="Admin")))
    ws = wb.active
    assert ws.sheet_view.rightToLeft
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "تقرير تجربة" in values and "سطر الفوتر الأول" in values
    assert "سطر الفوتر الأول" in ws.oddFooter.center.text


def test_template_export_has_no_footer_rows(db):
    settings = svc.get_settings()
    settings.footer_text = "فوتر"
    df = pd.DataFrame(columns=["كود", "اسم"])
    ws = load_workbook(io.BytesIO(to_excel(df, "قالب", settings, include_footer=False))).active
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "فوتر" not in values


def test_downloaded_template_reads_back(db):
    from core.services.accounts import import_template
    from ui.common import _read_template

    tpl = import_template()
    filled = pd.DataFrame([["6115", "نظافة", "", "61"]], columns=tpl.columns)
    data = to_excel(filled, "قالب استيراد", svc.get_settings(), include_footer=False)
    back = _read_template(io.BytesIO(data), list(tpl.columns))
    assert back.iloc[0]["كود الحساب"] == "6115" and len(back) == 1
