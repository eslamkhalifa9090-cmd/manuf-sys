from core import config
from core.services import system


def test_neon_style_urls_get_the_psycopg_driver():
    assert config.normalize_db_url("postgres://u:p@h/db?sslmode=require") == "postgresql+psycopg://u:p@h/db?sslmode=require"
    assert config.normalize_db_url("postgresql://u:p@h/db") == "postgresql+psycopg://u:p@h/db"
    assert config.normalize_db_url("postgresql+psycopg://u@h/db") == "postgresql+psycopg://u@h/db"
    assert config.normalize_db_url("sqlite:///x.db") == "sqlite:///x.db"


def test_env_var_wins(monkeypatch):
    monkeypatch.setenv("MANUF_DB_URL", "postgres://a@b/c")
    assert config.db_url() == "postgresql+psycopg://a@b/c"


def test_excel_dump_has_every_table(db, master):
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(system.excel_dump()), read_only=True)
    assert {"users", "accounts", "company_settings", "audit_log"} <= set(wb.sheetnames)
