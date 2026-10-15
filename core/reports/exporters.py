"""Excel and PDF export for any table, with company header, logo and footer.

Both exporters take a DataFrame plus the company settings, so every screen and
report gets the same layout. PDF text is Arabic-shaped and laid out right to left.
"""
import io
from datetime import date, datetime
from numbers import Number

import arabic_reshaper
import pandas as pd
from bidi.algorithm import get_display
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader, simpleSplit
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from core.config import FONTS_DIR

HEADER_FILL = "1F4E79"
ZEBRA_FILL = "EEF3F8"
FONT = "Amiri"
FONT_BOLD = "Amiri-Bold"
_fonts_registered = False


def _register_fonts() -> None:
    global _fonts_registered
    if not _fonts_registered:
        pdfmetrics.registerFont(TTFont(FONT, str(FONTS_DIR / "Amiri-Regular.ttf")))
        pdfmetrics.registerFont(TTFont(FONT_BOLD, str(FONTS_DIR / "Amiri-Bold.ttf")))
        _fonts_registered = True


def ar(text) -> str:
    """Shape and reorder Arabic text for drawing left-to-right."""
    return get_display(arabic_reshaper.reshape(str(text)))


def format_value(value, decimals: int = 2) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, bool):
        return "نعم" if value else ""
    if isinstance(value, Number):
        if float(value).is_integer() and not isinstance(value, float):
            return f"{int(value):,}"
        return f"{value:,.{decimals}f}"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M") if (value.hour or value.minute) else value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value)


def _signatures(settings) -> list[str]:
    return [s.strip() for s in (settings.signature_labels or "").split(",") if s.strip()]


def _footer_lines(settings, printed_by: str | None) -> list[str]:
    lines = []
    if settings.footer_text:
        lines.extend(l for l in settings.footer_text.splitlines() if l.strip())
    meta = []
    if settings.show_print_date:
        meta.append(f"تاريخ الطباعة: {datetime.now():%Y-%m-%d %H:%M}")
    if settings.show_printed_by and printed_by:
        meta.append(f"طبع بواسطة: {printed_by}")
    if meta:
        lines.append("   |   ".join(meta))
    return lines


# ---------------------------------------------------------------- Excel

def to_excel(df: pd.DataFrame, title: str, settings, subtitle: str | None = None,
             printed_by: str | None = None, sheet_name: str = "تقرير", include_footer: bool = True,
             totals: dict | None = None, bold_rows: set[int] | None = None) -> bytes:
    """`bold_rows` holds 0-based positions of rows to print bold and shaded, such as statement subtotals."""
    bold_rows = bold_rows or set()
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    ws.sheet_view.rightToLeft = True
    ncols = max(len(df.columns), 1)
    last_col = get_column_letter(ncols)
    thin = Side(style="thin", color="B7C4D1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    header_rows = [(settings.name_ar, 16, True), (title, 13, True)]
    if subtitle:
        header_rows.append((subtitle, 11, False))
    first_col = 1
    if settings.logo:
        try:
            img = XLImage(io.BytesIO(settings.logo))
            ratio = 60 / max(img.height, 1)
            img.height, img.width = 60, int(img.width * ratio)
            ws.add_image(img, "A1")
            first_col = 2 if ncols > 2 else 1
        except Exception:  # an unreadable logo must not block the export
            pass
    for i, (text, size, bold) in enumerate(header_rows, start=1):
        start = get_column_letter(first_col)
        ws.merge_cells(f"{start}{i}:{last_col}{i}" if ncols > first_col else f"{start}{i}:{start}{i}")
        cell = ws[f"{start}{i}"]
        cell.value = text
        cell.font = Font(bold=bold, size=size, color=HEADER_FILL if i == 1 else "000000")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    start_row = len(header_rows) + 2

    for j, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=j, value=str(col))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    decimals = settings.decimals if settings.decimals is not None else 2
    num_format = "#,##0" + ("." + "0" * decimals if decimals else "")
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, value in enumerate(row, start=1):
            if isinstance(value, bool):
                value = "نعم" if value else ""
            elif value is not None and isinstance(value, float) and pd.isna(value):
                value = None
            elif isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            cell = ws.cell(row=i, column=j, value=value)
            cell.border = border
            if isinstance(value, float):
                cell.number_format = num_format
            elif isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd"
            if i - start_row - 1 in bold_rows:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="DDEBF7")
            elif (i - start_row) % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=ZEBRA_FILL)

    for j, col in enumerate(df.columns, start=1):
        lengths = [len(str(col))] + [len(format_value(v, decimals)) for v in df.iloc[:200, j - 1]]
        ws.column_dimensions[get_column_letter(j)].width = min(max(lengths) + 3, 60)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    if len(df):
        ws.auto_filter.ref = f"A{start_row}:{last_col}{start_row + len(df)}"

    row = start_row + len(df) + 2
    if totals:
        trow = start_row + len(df) + 1
        for j, col in enumerate(df.columns, start=1):
            value = totals.get(col, "الإجمالي" if j == 1 else None)
            cell = ws.cell(row=trow, column=j, value=value)
            cell.font = Font(bold=True)
            cell.border = border
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
            if isinstance(value, float):
                cell.number_format = num_format
        row += 1
    for line in _footer_lines(settings, printed_by) if include_footer else []:
        ws.cell(row=row, column=1, value=line).font = Font(italic=True, size=9, color="555555")
        row += 1
    if settings.footer_text:
        ws.oddFooter.center.text = settings.footer_text.replace("\n", " ")
    if settings.show_page_numbers:
        ws.oddFooter.left.text = "صفحة &P من &N"
    ws.page_setup.orientation = "landscape" if ncols > 7 else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"{start_row}:{start_row}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF

def _rtl_paragraph(text: str, style: ParagraphStyle, width: float) -> Paragraph:
    """Wrap in logical order first, then reorder each line, so multi-line cells read top to bottom."""
    shaped = arabic_reshaper.reshape(str(text))
    lines = simpleSplit(shaped, style.fontName, style.fontSize, max(width, 10)) or [""]
    visual = "<br/>".join(get_display(l).replace("&", "&amp;").replace("<", "&lt;") for l in lines)
    return Paragraph(visual, style)


def _make_canvas(settings, printed_by, pagesize):
    footer = _footer_lines(settings, printed_by)

    class FooterCanvas(rl_canvas.Canvas):
        """Two-pass canvas so each page can print 'page x of y'."""

        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._pages = []

        def showPage(self):
            self._pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._pages)
            for state in self._pages:
                self.__dict__.update(state)
                self._draw_footer(total)
                super().showPage()
            super().save()

        def _draw_footer(self, total):
            w, _ = pagesize
            self.setStrokeColor(colors.HexColor("#" + HEADER_FILL))
            self.setLineWidth(0.6)
            self.line(15 * mm, 16 * mm, w - 15 * mm, 16 * mm)
            self.setFont(FONT, 8.5)
            self.setFillColor(colors.HexColor("#444444"))
            y = 12 * mm
            for line in footer[:3]:
                self.drawCentredString(w / 2, y, ar(line))
                y -= 4 * mm
            if settings.show_page_numbers:
                self.drawString(15 * mm, 12 * mm, ar(f"صفحة {self._pageNumber} من {total}"))

    return FooterCanvas


def to_pdf(df: pd.DataFrame, title: str, settings, subtitle: str | None = None,
           printed_by: str | None = None, landscape_mode: bool | None = None,
           totals: dict | None = None, bold_rows: set[int] | None = None) -> bytes:
    """Render a table to PDF. `totals` maps column name -> value for a final bold row;
    `bold_rows` holds 0-based positions of rows to print bold and shaded."""
    bold_rows = bold_rows or set()
    _register_fonts()
    if landscape_mode is None:
        landscape_mode = len(df.columns) > 6
    pagesize = landscape(A4) if landscape_mode else A4
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=pagesize, rightMargin=15 * mm, leftMargin=15 * mm,
                            topMargin=12 * mm, bottomMargin=24 * mm, title=title)
    avail = pagesize[0] - 30 * mm
    decimals = settings.decimals if settings.decimals is not None else 2

    h1 = ParagraphStyle("h1", fontName=FONT_BOLD, fontSize=16, leading=20, alignment=TA_RIGHT,
                        textColor=colors.HexColor("#" + HEADER_FILL))
    h2 = ParagraphStyle("h2", fontName=FONT_BOLD, fontSize=13, leading=17, alignment=TA_CENTER)
    h3 = ParagraphStyle("h3", fontName=FONT, fontSize=10, leading=13, alignment=TA_CENTER,
                        textColor=colors.HexColor("#555555"))
    cell = ParagraphStyle("cell", fontName=FONT, fontSize=8.5, leading=11, alignment=TA_RIGHT)
    head = ParagraphStyle("head", parent=cell, fontName=FONT_BOLD, textColor=colors.white, alignment=TA_CENTER)
    num = ParagraphStyle("num", parent=cell, alignment=TA_CENTER)

    story = []
    # Header: company block on the right, logo on the left.
    company = [_rtl_paragraph(settings.name_ar, h1, avail * 0.7)]
    small = ParagraphStyle("small", parent=cell, fontSize=8.5, textColor=colors.HexColor("#555555"))
    details = " - ".join(x for x in (settings.activity, settings.address, settings.phone) if x)
    if details:
        company.append(_rtl_paragraph(details, small, avail * 0.7))
    if settings.tax_card or settings.commercial_register:
        ids = []
        if settings.commercial_register:
            ids.append(f"س.ت: {settings.commercial_register}")
        if settings.tax_card:
            ids.append(f"ب.ض: {settings.tax_card}")
        company.append(_rtl_paragraph(" | ".join(ids), small, avail * 0.7))
    logo_cell = ""
    if settings.logo:
        try:
            reader = ImageReader(io.BytesIO(settings.logo))
            iw, ih = reader.getSize()
            h = 20 * mm
            from reportlab.platypus import Image as RLImage

            logo_cell = RLImage(io.BytesIO(settings.logo), width=h * iw / ih, height=h)
        except Exception:
            logo_cell = ""
    header = Table([[logo_cell, company]], colWidths=[avail * 0.3, avail * 0.7])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (0, 0), "LEFT")]))
    story += [header, Spacer(1, 3 * mm), _rtl_paragraph(title, h2, avail)]
    if subtitle:
        story.append(_rtl_paragraph(subtitle, h3, avail))
    story.append(Spacer(1, 4 * mm))

    # Table: columns reversed so the first column sits on the right.
    columns = list(df.columns)[::-1]
    sample = df.head(300)
    widths = []
    for col in columns:
        # Wide enough for the longest header word and the longest value, within limits.
        head_w = max(pdfmetrics.stringWidth(ar(w), FONT_BOLD, 8.5) for w in (str(col).split() or [""]))
        vals = [pdfmetrics.stringWidth(ar(format_value(v, decimals)), FONT, 8.5) for v in sample[col]]
        widths.append(min(max([head_w] + vals), 220) + 8)
    scale = avail / (sum(widths) or 1)
    widths = [w * scale for w in widths]
    numeric = {c: pd.api.types.is_numeric_dtype(df[c]) and not pd.api.types.is_bool_dtype(df[c]) for c in columns}

    data = [[_rtl_paragraph(c, head, w - 4) for c, w in zip(columns, widths)]]
    cell_bold = ParagraphStyle("cellb", parent=cell, fontName=FONT_BOLD)
    num_bold = ParagraphStyle("numb", parent=num, fontName=FONT_BOLD)
    for pos, (_, rec) in enumerate(df[columns].iterrows()):
        strong = pos in bold_rows
        data.append([
            _rtl_paragraph(format_value(rec[c], decimals),
                           (num_bold if strong else num) if numeric[c] else (cell_bold if strong else cell), w - 4)
            for c, w in zip(columns, widths)
        ])
    if totals:
        bold = ParagraphStyle("tot", parent=num, fontName=FONT_BOLD)
        row = []
        for i, (c, w) in enumerate(zip(columns, widths)):
            if c in totals:
                row.append(_rtl_paragraph(format_value(totals[c], decimals), bold, w - 4))
            elif i == len(columns) - 1:
                row.append(_rtl_paragraph("الإجمالي", bold, w - 4))
            else:
                row.append("")
        data.append(row)

    table = Table(data, colWidths=widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + HEADER_FILL)),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B7C4D1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]
    for r in range(2, len(data), 2):
        style.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#" + ZEBRA_FILL)))
    for pos in bold_rows:
        if pos < len(df):
            style.append(("BACKGROUND", (0, pos + 1), (-1, pos + 1), colors.HexColor("#DDEBF7")))
    if totals:
        style.append(("BACKGROUND", (0, len(data) - 1), (-1, len(data) - 1), colors.HexColor("#D9E2EC")))
    table.setStyle(TableStyle(style))
    story.append(table)

    sigs = _signatures(settings)
    if sigs:
        story.append(Spacer(1, 14 * mm))
        sig_style = ParagraphStyle("sig", parent=cell, fontName=FONT_BOLD, fontSize=10, alignment=TA_CENTER)
        sig_w = avail / len(sigs)
        row = [_rtl_paragraph(s, sig_style, sig_w - 4) for s in sigs[::-1]]
        line = ["...................................."] * len(sigs)
        sig_table = Table([row, line], colWidths=[sig_w] * len(sigs))
        sig_table.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER")]))
        story.append(sig_table)

    doc.build(story, canvasmaker=_make_canvas(settings, printed_by, pagesize))
    return buf.getvalue()
