"""Opening balances from one Excel workbook: accounts, customers and suppliers, stock, fixed assets and molds.

Everything lands at the cut-over date as pending documents: one opening journal (accounts, party balances,
asset and mold cost with their accumulated depreciation) and one opening stock voucher. The registers get
the same figures as their opening balances, so the books and the registers start equal. The journal is
balanced on 3303 and the stock voucher credits 3303, so 3303 comes back to zero when the opening is complete.
"""
import io
from dataclasses import dataclass, field, replace
from datetime import date, datetime, timedelta

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select

from core.db import session_scope
from core.models import Account, Item, ItemGroup, Party
from core.models.assets import ADMIN, CATEGORY_KEYS, FACTORY, FixedAsset
from core.models.documents import APPROVED, JOURNAL_OPENING
from core.models.inventory import STOCK_OPENING
from core.models.molds import COUNTER_UNITS, CUSTOMER_OWNED, FACTORY as MOLD_FACTORY, Mold
from core.models.parties import BOTH, CUSTOMER, PARTY_KINDS, SUPPLIER
from core.services import audit, fixed_assets, items, journals, molds, parties, stock_vouchers
from core.services.approval import r2
from core.services.errors import BusinessError
from core.services.journals import OPENING_SCREEN
from core.services.periods import is_closed

READY = "جاهز"
NEW, EXISTING = "جديد", "موجود"

SHEETS = {
    "accounts": ("أرصدة الحسابات", ["كود الحساب", "اسم الحساب", "مدين", "دائن", "البيان"]),
    "parties": ("العملاء والموردين", ["الاسم", "النوع", "مدين", "دائن", "التليفون", "العنوان"]),
    "stock": ("المخزون", ["كود الصنف", "اسم الصنف", "المجموعة", "الوحدة", "الكمية", "تكلفة الوحدة"]),
    "assets": ("الأصول الثابتة", ["كود الأصل", "اسم الأصل", "كود حساب الفئة", "القسم", "تاريخ الشراء", "التكلفة",
                                   "الخردة", "العمر بالشهور", "مجمع الإهلاك"]),
    "molds": ("الأسطمبات", ["كود الأسطمبة", "اسم الأسطمبة", "الملكية", "العميل", "وحدة العداد", "التجاويف",
                            "العمر الافتراضي", "الخردة", "التكلفة", "مجمع الاستهلاك", "المستهلك قبل البداية"]),
}
DEPARTMENTS = {"المصنع": FACTORY, "الإدارة": ADMIN}
MOLD_OWNERSHIP = {"ملك المصنع": MOLD_FACTORY, "ملك العميل": CUSTOMER_OWNED}
UNITS_BY_LABEL = {v: k for k, v in COUNTER_UNITS.items()}
PARTY_BY_LABEL = {v: k for k, v in PARTY_KINDS.items()}
# Accounts filled from another sheet, so their figure always matches its register or subledger.
FROM_SHEET = {"ar": "العملاء والموردين", "ap": "العملاء والموردين", "mold_asset": "الأسطمبات",
              "mold_acc_amortization": "الأسطمبات", "opening_balance_equity": "(بيتحسب تلقائي)"}

INSTRUCTIONS = [
    "املأ الشيتات اللي عندك بس، وسيب الباقي فاضي. كل الأرقام في تاريخ بداية التشغيل.",
    "أرصدة الحسابات: كود الحساب الفرعي والرصيد في مدين أو دائن. حسابات العملاء والموردين والمخزون والأصول "
    "والأسطمبات بتتسجل من شيتاتها عشان السجلات تطابق الدفاتر.",
    "العملاء والموردين: النوع عميل / مورد / عميل ومورد. الرصيد المدين على 1401 والدائن على 2101 للنوع «عميل ومورد».",
    "المخزون: كود الصنف أو اسمه. لو الصنف مش مسجل هيتسجل بالمجموعة والوحدة المكتوبين.",
    "الأصول الثابتة: كود حساب الفئة (1101 ماكينات، 1102 عدد وأدوات، ...)، القسم المصنع أو الإدارة، "
    "ومجمع الإهلاك لحد تاريخ البداية. الإهلاك الشهري بيبدأ من الشهر اللي بعده.",
    "الأسطمبات: الملكية ملك المصنع أو ملك العميل، وحدة العداد قطعة أو جوز، والعمر والمستهلك بوحدة العداد.",
    "الفرق بين المدين والدائن بيتقفل على 3303 أرصدة افتتاحية؛ لو الأرصدة كاملة ومتزنة الفرق ده بيرجع صفر "
    "بعد اعتماد إذن المخزن.",
]


# ---------------------------------------------------------------- template

def _write_sheet(ws, title: str, columns: list[str], rows: list[list] | None = None) -> None:
    ws.title = title
    ws.sheet_view.rightToLeft = True
    ws.append(columns)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center")
    for r in rows or []:
        ws.append(r)
    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, len(col) + 6)
    ws.freeze_panes = "A2"


def template(frames: dict[str, pd.DataFrame] | None = None) -> bytes:
    """The opening workbook: an instructions sheet then one sheet per kind, optionally pre-filled."""
    wb = Workbook()
    ws = wb.active
    ws.title = "تعليمات"
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 140
    ws.append(["استيراد الأرصدة الافتتاحية"])
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    for line in INSTRUCTIONS:
        ws.append([line])
    for key, (title, columns) in SHEETS.items():
        rows = None
        if frames is not None and key in frames:
            rows = frames[key].reindex(columns=columns).astype(object).where(frames[key].notna(), None).values.tolist()
        _write_sheet(wb.create_sheet(), title, columns, rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_workbook(data: bytes) -> dict[str, pd.DataFrame]:
    """One DataFrame per known sheet; missing sheets come back empty."""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise BusinessError("تعذر قراءة الملف. استخدم القالب اللي اتحمل من البرنامج.") from exc
    out = {}
    for key, (title, columns) in SHEETS.items():
        if title not in wb.sheetnames:
            out[key] = pd.DataFrame(columns=columns)
            continue
        rows = list(wb[title].iter_rows(values_only=True))
        header = [str(h).strip() if h is not None else "" for h in (rows[0] if rows else [])]
        missing = [c for c in columns if c not in header]
        if missing:
            raise BusinessError(f"شيت «{title}»: أعمدة ناقصة: " + "، ".join(missing))
        df = pd.DataFrame([r for r in rows[1:] if any(v not in (None, "") for v in r)],
                          columns=header) if len(rows) > 1 else pd.DataFrame(columns=header)
        out[key] = df.reindex(columns=columns)
    return out


# ---------------------------------------------------------------- validation

def _text(v) -> str:
    if v is None or (isinstance(v, float) and v != v):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _num(v) -> float | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return 0.0
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return 0.0 if f != f else f


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if v in (None, ""):
        return None
    try:
        return pd.to_datetime(v).date()
    except Exception:
        return None


@dataclass
class Plan:
    """What the import would do; each checked frame carries a «الحالة» column."""
    as_of: date
    checked: dict[str, pd.DataFrame] = field(default_factory=dict)
    journal_lines: list[dict] = field(default_factory=list)
    stock_lines: list[dict] = field(default_factory=list)
    errors: int = 0
    warnings: list[str] = field(default_factory=list)

    @property
    def debit(self) -> float:
        return r2(sum(float(l["debit"]) for l in self.journal_lines))

    @property
    def credit(self) -> float:
        return r2(sum(float(l["credit"]) for l in self.journal_lines))

    @property
    def stock_value(self) -> float:
        return r2(sum(l["qty"] * l["unit_cost"] for l in self.stock_lines))

    @property
    def equity_difference(self) -> float:
        """What 3303 holds once both documents are approved: zero when the opening is complete and balanced."""
        return r2(self.debit - self.credit + self.stock_value)


def _mark(df: pd.DataFrame, status: list[str], note: list[str] | None = None) -> pd.DataFrame:
    df = df.copy()
    df["الحالة"] = status
    if note is not None:
        df["ملاحظة"] = note
    return df


def plan(frames: dict[str, pd.DataFrame], as_of: date) -> Plan:
    if as_of is None:
        raise BusinessError("اختار تاريخ بداية التشغيل")
    p = Plan(as_of)
    with session_scope() as s:
        if is_closed(s, as_of):
            raise BusinessError("فترة تاريخ البداية مقفولة")
        accounts = {a.code: a for a in s.scalars(select(Account))}
        by_key = {a.system_key: a for a in accounts.values() if a.system_key}
        inventory_ids = set(s.scalars(select(ItemGroup.inventory_account_id)))
        category_ids = {by_key[k].id for k in CATEGORY_KEYS if k in by_key} | \
            {by_key[v].id for v in CATEGORY_KEYS.values() if v in by_key}
        party_by_name = {x.name.strip().lower(): x for x in s.scalars(select(Party))}
        items_by_code = {i.code: i for i in s.scalars(select(Item)) if i.code}
        items_by_name = {i.name.strip().lower(): i for i in s.scalars(select(Item))}
        groups = {g.name: g for g in s.scalars(select(ItemGroup).where(ItemGroup.is_active))}
        asset_names = {a.name.strip().lower() for a in s.scalars(select(FixedAsset))}
        asset_codes = set(s.scalars(select(FixedAsset.code)))
        mold_names = {m.name.strip().lower() for m in s.scalars(select(Mold))}
        mold_codes = set(s.scalars(select(Mold.code)))
        if s.scalar(select(journals.ManualJournal.id).where(journals.ManualJournal.kind == JOURNAL_OPENING,
                                                            journals.ManualJournal.status == APPROVED).limit(1)):
            p.warnings.append("فيه قيد أرصدة افتتاحية معتمد قبل كده؛ الاستيراد هيضيف قيد تاني. راجعه قبل الاعتماد.")
    equity = by_key.get("opening_balance_equity")

    # Accounts
    df, status = frames.get("accounts", pd.DataFrame(columns=SHEETS["accounts"][1])), []
    for r in df.to_dict("records"):
        code, dr, cr = _text(r["كود الحساب"]).removesuffix(".0"), _num(r["مدين"]), _num(r["دائن"])
        acc = accounts.get(code)
        if acc is None:
            status.append("الحساب غير موجود")
        elif not acc.is_postable:
            status.append("حساب رئيسي؛ اكتب الحساب الفرعي")
        elif acc.system_key in FROM_SHEET:
            status.append(f"بيتسجل من شيت {FROM_SHEET[acc.system_key]}")
        elif acc.id in inventory_ids:
            status.append("حساب مخزون؛ بيتسجل من شيت المخزون")
        elif acc.id in category_ids:
            status.append("حساب أصول ثابتة؛ بيتسجل من شيت الأصول الثابتة")
        elif dr is None or cr is None or dr < 0 or cr < 0:
            status.append("المبلغ لازم يكون رقم موجب")
        elif (dr > 0) == (cr > 0):
            status.append("اكتب مدين أو دائن (واحد بس)")
        else:
            status.append(READY)
            p.journal_lines.append({"account_id": acc.id, "debit": r2(dr), "credit": r2(cr),
                                    "memo": _text(r["البيان"]) or "رصيد افتتاحي"})
    p.checked["accounts"] = _mark(df, status)

    # Customers and suppliers
    df, status, note, seen = frames.get("parties", pd.DataFrame(columns=SHEETS["parties"][1])), [], [], set()
    for r in df.to_dict("records"):
        name, kind_label = _text(r["الاسم"]), _text(r["النوع"]) or PARTY_KINDS[CUSTOMER]
        dr, cr = _num(r["مدين"]), _num(r["دائن"])
        existing = party_by_name.get(name.lower())
        kind = PARTY_BY_LABEL.get(kind_label)
        if not name:
            status.append("الاسم مطلوب")
        elif name.lower() in seen:
            status.append("الاسم متكرر في الشيت")
        elif kind is None:
            status.append("النوع لازم يكون: " + " / ".join(PARTY_KINDS.values()))
        elif existing is not None and existing.kind not in (kind, BOTH):
            status.append(f"مسجل كـ{PARTY_KINDS[existing.kind]}")
        elif dr is None or cr is None or dr < 0 or cr < 0:
            status.append("المبلغ لازم يكون رقم موجب")
        elif dr and cr:
            status.append("اكتب مدين أو دائن (واحد بس)")
        else:
            status.append(READY)
            seen.add(name.lower())
            if dr or cr:
                key = "ap" if kind == SUPPLIER or (kind == BOTH and cr) else "ar"
                p.journal_lines.append({"account_id": by_key[key].id, "debit": r2(dr), "credit": r2(cr),
                                        "memo": f"رصيد افتتاحي {name}", "party_name": name, "party_kind": kind})
        note.append(EXISTING if existing is not None else NEW)
    p.checked["parties"] = _mark(df, status, note)

    # Stock
    df, status, note, seen = frames.get("stock", pd.DataFrame(columns=SHEETS["stock"][1])), [], [], set()
    for r in df.to_dict("records"):
        code, name = _text(r["كود الصنف"]), _text(r["اسم الصنف"])
        group, unit = _text(r["المجموعة"]), _text(r["الوحدة"])
        qty, cost = _num(r["الكمية"]), _num(r["تكلفة الوحدة"])
        item = items_by_code.get(code) if code else None
        item = item or (items_by_name.get(name.lower()) if name else None)
        ident = (item.id if item else None) or (code or name).lower()
        if not (code or name):
            status.append("اكتب كود الصنف أو اسمه")
        elif ident in seen:
            status.append("الصنف متكرر في الشيت")
        elif item is None and not name:
            status.append("الكود مش مسجل؛ اكتب اسم الصنف عشان يتسجل")
        elif item is None and group not in groups:
            status.append("صنف جديد: المجموعة غير موجودة")
        elif item is None and not unit:
            status.append("صنف جديد: الوحدة مطلوبة")
        elif qty is None or cost is None or qty <= 0 or cost < 0:
            status.append("الكمية لازم تكون أكبر من صفر والتكلفة رقم موجب")
        else:
            status.append(READY)
            seen.add(ident)
            p.stock_lines.append({"item_id": item.id if item else None, "code": code, "name": name, "group": group,
                                  "unit": unit, "qty": float(qty), "unit_cost": round(float(cost), 6)})
        note.append(EXISTING if item is not None else NEW)
    p.checked["stock"] = _mark(df, status, note)

    # Fixed assets
    df, status, seen = frames.get("assets", pd.DataFrame(columns=SHEETS["assets"][1])), [], set()
    for r in df.to_dict("records"):
        code, name = _text(r["كود الأصل"]), _text(r["اسم الأصل"])
        cat = accounts.get(_text(r["كود حساب الفئة"]).removesuffix(".0"))
        dept = DEPARTMENTS.get(_text(r["القسم"]) or "المصنع")
        acquired = _date(r["تاريخ الشراء"])
        cost, salvage, life, acc = (_num(r["التكلفة"]), _num(r["الخردة"]), _num(r["العمر بالشهور"]),
                                    _num(r["مجمع الإهلاك"]))
        if not name:
            status.append("اسم الأصل مطلوب")
        elif name.lower() in asset_names or name.lower() in seen:
            status.append("الأصل مسجل قبل كده")
        elif code and code in asset_codes:
            status.append("الكود مستخدم")
        elif cat is None or cat.system_key not in CATEGORY_KEYS:
            status.append("كود حساب الفئة لازم يكون من حسابات الأصول (1101 - 1107 ماعدا الأسطمبات)")
        elif dept is None:
            status.append("القسم لازم يكون المصنع أو الإدارة")
        elif acquired is None or acquired > p.as_of:
            status.append("تاريخ الشراء مطلوب وقبل تاريخ البداية")
        elif None in (cost, salvage, life, acc) or cost <= 0 or life <= 0 or salvage < 0 or acc < 0:
            status.append("التكلفة والعمر أكبر من صفر، والخردة والمجمع أرقام موجبة")
        elif salvage >= cost:
            status.append("الخردة لازم تكون أقل من التكلفة")
        elif acc > cost - salvage + 0.005:
            status.append("المجمع أكبر من القيمة القابلة للإهلاك")
        else:
            status.append(READY)
            seen.add(name.lower())
            p.journal_lines.append({"account_id": cat.id, "debit": r2(cost), "credit": 0.0, "memo": f"تكلفة {name}"})
            if acc:
                p.journal_lines.append({"account_id": by_key[CATEGORY_KEYS[cat.system_key]].id, "debit": 0.0,
                                        "credit": r2(acc), "memo": f"مجمع إهلاك {name}"})
            p.checked.setdefault("_assets", []).append(  # type: ignore[arg-type]
                {"code": code, "name": name, "account_id": cat.id, "department": dept, "acquisition_date": acquired,
                 "cost": r2(cost), "salvage_value": r2(salvage), "life_months": int(life),
                 "opening_accumulated": r2(acc)})
    p.checked["assets"] = _mark(df, status)

    # Molds
    df, status, seen = frames.get("molds", pd.DataFrame(columns=SHEETS["molds"][1])), [], set()
    for r in df.to_dict("records"):
        code, name = _text(r["كود الأسطمبة"]), _text(r["اسم الأسطمبة"])
        ownership = MOLD_OWNERSHIP.get(_text(r["الملكية"]) or "ملك المصنع")
        customer = party_by_name.get(_text(r["العميل"]).lower()) if _text(r["العميل"]) else None
        customer_new = _text(r["العميل"]).lower() in {n.lower() for n in p.checked["parties"]["الاسم"].map(_text)}
        unit = UNITS_BY_LABEL.get(_text(r["وحدة العداد"]) or COUNTER_UNITS["piece"])
        cav, life, salvage, cost, acc, used = (_num(r["التجاويف"]), _num(r["العمر الافتراضي"]), _num(r["الخردة"]),
                                               _num(r["التكلفة"]), _num(r["مجمع الاستهلاك"]),
                                               _num(r["المستهلك قبل البداية"]))
        if not name:
            status.append("اسم الأسطمبة مطلوب")
        elif name.lower() in mold_names or name.lower() in seen:
            status.append("الأسطمبة مسجلة قبل كده")
        elif code and code in mold_codes:
            status.append("الكود مستخدم")
        elif ownership is None:
            status.append("الملكية لازم تكون ملك المصنع أو ملك العميل")
        elif unit is None:
            status.append("وحدة العداد لازم تكون قطعة أو جوز")
        elif None in (cav, life, salvage, cost, acc, used) or cav < 1 or min(life, salvage, cost, acc, used) < 0:
            status.append("التجاويف 1 أو أكتر، وباقي الأرقام موجبة")
        elif ownership == CUSTOMER_OWNED and customer is None and not customer_new:
            status.append("اكتب اسم العميل صاحب الأسطمبة زي ما هو مسجل")
        elif ownership == CUSTOMER_OWNED and (cost or acc):
            status.append("أسطمبة العميل ملهاش تكلفة ولا استهلاك")
        elif ownership == MOLD_FACTORY and life <= 0:
            status.append("العمر الافتراضي مطلوب لأسطمبة المصنع")
        elif ownership == MOLD_FACTORY and used >= life:
            status.append("المستهلك لازم يكون أقل من العمر الافتراضي")
        elif acc > cost + 0.005:
            status.append("مجمع الاستهلاك أكبر من التكلفة")
        else:
            status.append(READY)
            seen.add(name.lower())
            if cost:
                p.journal_lines.append({"account_id": by_key["mold_asset"].id, "debit": r2(cost), "credit": 0.0,
                                        "memo": f"تكلفة أسطمبة {name}"})
            if acc:
                p.journal_lines.append({"account_id": by_key["mold_acc_amortization"].id, "debit": 0.0,
                                        "credit": r2(acc), "memo": f"مجمع استهلاك أسطمبة {name}"})
            p.checked.setdefault("_molds", []).append(  # type: ignore[arg-type]
                {"code": code, "name": name, "ownership": ownership, "customer": _text(r["العميل"]),
                 "counter_unit": unit, "cavities": int(cav), "useful_life": life, "salvage_value": salvage,
                 "opening_cost": cost, "opening_amortization": acc, "opening_usage": used})
    p.checked["molds"] = _mark(df, status)

    p.errors = sum(int((v["الحالة"] != READY).sum()) for k, v in p.checked.items() if not k.startswith("_"))
    if equity is None:
        raise BusinessError("حساب 3303 أرصدة افتتاحية غير موجود في الشجرة")
    if not p.journal_lines and not p.stock_lines:
        p.warnings.append("مفيش أرصدة في الملف")
    elif p.equity_difference:
        side = "دائن" if p.equity_difference > 0 else "مدين"
        p.warnings.append(f"الأرصدة مش متزنة: هيفضل على 3303 رصيد {side} {abs(p.equity_difference):,.2f} بعد الاعتماد. "
                          "لو ده مقصود (مثلاً رأس المال مش مكتوب) كمّل، وإلا راجع الملف.")
    return p


# ---------------------------------------------------------------- import

def _as_opening_user(actor):
    """The importer creates the register records the opening needs, under the opening-balances permission."""
    extra = {screen: {"add"} for screen in (parties.SCREEN, items.SCREEN, fixed_assets.SCREEN, molds.SCREEN)}
    return replace(actor, perms={**actor.perms, **{k: actor.perms.get(k, set()) | v for k, v in extra.items()}})


def run(actor, frames: dict[str, pd.DataFrame], as_of: date, memo: str = "") -> dict:
    """Create the masters and the two pending documents. Refuses while any row is invalid."""
    if not actor.can(OPENING_SCREEN, "import") or not actor.can(OPENING_SCREEN, "add"):
        raise BusinessError("ليس لديك صلاحية استيراد الأرصدة الافتتاحية")
    p = plan(frames, as_of)
    if p.errors:
        raise BusinessError(f"فيه {p.errors} سطر غير صالح. صححهم في الملف وارفعه تاني.")
    if not p.journal_lines and not p.stock_lines:
        raise BusinessError("مفيش أرصدة في الملف")
    user = _as_opening_user(actor)
    memo = (memo or "").strip() or f"أرصدة افتتاحية في {as_of}"
    created = {"parties": 0, "items": 0, "assets": 0, "molds": 0}

    # Parties first: journal lines and customer-owned molds point at them.
    df = frames.get("parties", pd.DataFrame(columns=SHEETS["parties"][1]))
    with session_scope() as s:
        party_ids = {x.name.strip().lower(): x.id for x in s.scalars(select(Party))}
    for r in df.to_dict("records"):
        name = _text(r["الاسم"])
        if name and name.lower() not in party_ids:
            party_ids[name.lower()] = parties.save_party(user, None, name, PARTY_BY_LABEL[_text(r["النوع"]) or
                                                                                         PARTY_KINDS[CUSTOMER]],
                                                         _text(r["التليفون"]), _text(r["العنوان"]))
            created["parties"] += 1

    stock_lines = []
    with session_scope() as s:
        groups = {g.name: g.id for g in s.scalars(select(ItemGroup))}
    for l in p.stock_lines:
        item_id = l["item_id"]
        if item_id is None:
            item_id = items.save_item(user, None, l["name"], groups[l["group"]], l["unit"], l["code"])
            created["items"] += 1
        stock_lines.append({"item_id": item_id, "qty": l["qty"], "unit_cost": l["unit_cost"], "memo": ""})

    for a in p.checked.get("_assets", []):
        fixed_assets.save_asset(user, None, a | {"depreciation_start": a["acquisition_date"], "opening_date": as_of})
        created["assets"] += 1
    for m in p.checked.get("_molds", []):
        customer = party_ids.get(m["customer"].lower()) if m["customer"] else None
        molds.save_mold(user, None, {k: v for k, v in m.items() if k != "customer"}
                        | {"customer_id": customer, "opening_date": as_of})
        created["molds"] += 1

    journal_id = voucher_id = None
    if p.journal_lines:
        lines = [{"account_id": l["account_id"], "debit": l["debit"], "credit": l["credit"], "memo": l["memo"],
                  "party_id": party_ids.get(l["party_name"].lower()) if l.get("party_name") else None}
                 for l in p.journal_lines]
        journal_id = journals.save_journal(actor, None, as_of, memo, JOURNAL_OPENING, lines, auto_balance=True)
    if stock_lines:
        voucher_id = stock_vouchers.save_voucher(actor, None, STOCK_OPENING, as_of, stock_lines,
                                                 memo=f"{memo} - المخزون")
    with session_scope() as s:
        audit.log(s, actor, "import", "opening_balances", journal_id or voucher_id or 0,
                  {"as_of": as_of.isoformat(), **created, "journal": journal_id, "stock_voucher": voucher_id})
    return {"journal_id": journal_id, "voucher_id": voucher_id, "created": created,
            "equity_difference": p.equity_difference}


def day_before(d: date) -> date:
    return d - timedelta(days=1)
