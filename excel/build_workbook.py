"""Build the New Star Excel edition (phase 1): sheets, tables, lists and the opening data.

The VBA lives in excel/vba (modNewStar.bas + clsWbEvents.cls) and is imported once in Excel.
Usage: python excel/build_workbook.py <opening_balances.xlsx> <original_workbook.pkl> <out.xlsx> [<notes.md>]
(the pickle is the original workbook's sheets as row lists, see dump_workbook in import_history)
"""
import sys
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo

sys.path.insert(0, __file__.rsplit("/excel/", 1)[0])
sys.path.insert(0, __file__.rsplit("/", 1)[0])
from core.seed.coa import ACCOUNTS  # noqa: E402
from import_history import History  # noqa: E402

OPENING_DATE = date(2026, 3, 31)
NAVY, WHITE, INPUT, LIGHT = "1F4E78", "FFFFFF", "FFF2CC", "DDEBF7"
FONT = "Arial"
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = "#,##0.00;(#,##0.00);-"

# key, sheet name, Master, Admin, Accountant, Storekeeper
SHEETS = [
    ("HOME", "الرئيسية", "E", "E", "E", "E"),
    ("JE_ENTRY", "قيد جديد", "E", "N", "E", "N"),
    ("VOUCHER", "سند قبض وصرف", "E", "N", "E", "N"),
    ("SALES_ENTRY", "فاتورة بيع", "E", "N", "E", "N"),
    ("STOCK_ENTRY", "إذن مخزن", "E", "N", "E", "E"),
    ("APPROVALS", "الاعتمادات", "E", "E", "N", "N"),
    ("JE", "دفتر اليومية", "E", "V", "V", "N"),
    ("SALES", "سجل المبيعات", "E", "V", "V", "N"),
    ("STOCK", "حركة المخزن", "E", "V", "V", "V"),
    ("TB", "ميزان المراجعة", "E", "V", "V", "N"),
    ("LEDGER", "كشف حساب", "E", "V", "V", "N"),
    ("PARTY_BAL", "أرصدة العملاء والموردين", "E", "V", "V", "N"),
    ("SALES_REP", "تقرير المبيعات", "E", "V", "V", "N"),
    ("ITEM_CARD", "كارت صنف", "E", "V", "V", "V"),
    ("FS", "القوائم المالية", "E", "V", "N", "N"),
    ("COA", "شجرة الحسابات", "E", "V", "E", "N"),
    ("PARTIES", "العملاء والموردين", "E", "V", "E", "N"),
    ("ITEMS", "الأصناف", "E", "V", "E", "E"),
    ("ASSETS", "الأصول الثابتة", "E", "V", "E", "N"),
    ("MOLDS", "الأسطمبات", "E", "V", "E", "V"),
    ("MOLD_USE", "تشغيل الأسطمبات", "E", "V", "E", "E"),
    ("LOG", "سجل الحركات", "E", "V", "N", "N"),
    ("USERS", "المستخدمين", "E", "N", "N", "N"),
    ("PERMS", "الصلاحيات", "E", "N", "N", "N"),
    ("SETTINGS", "الإعدادات", "E", "N", "N", "N"),
    ("LOGIN", "الدخول", "-", "-", "-", "-"),
    ("WARN", "تنبيه", "-", "-", "-", "-"),
    ("LISTS", "قوائم", "-", "-", "-", "-"),
]
NAME = {k: n for k, n, *_ in SHEETS}

ROLES = [("Master", "ماستر"), ("Admin", "مدير عام"), ("Accountant", "محاسب"), ("Storekeeper", "أمين مخزن")]
USERS = [("master", "ماستر النظام", "ماستر"), ("admin", "المدير العام", "مدير عام"),
         ("acc", "المحاسب", "محاسب"), ("store", "أمين المخزن", "أمين مخزن")]
DEFAULT_PASS = "1234"

GROUPS = [("خامات", "1301", "RM"), ("تضريبة وخامات وسيطة", "1302", "CM"), ("إنتاج تام", "1304", "FG"),
          ("قطع غيار ومهمات", "1305", "SP"), ("مستلزمات تعبئة", "1306", "PK"),
          ("عينات وهدايا واردة", "1307", "SM"), ("اكسسوارات", "1308", "AC")]

MSG = {
    "APP_TITLE": "New Star", "YES": "نعم", "NO": "لا", "WELCOME": "أهلاً", "TOTAL": "الإجمالي",
    "OPENING": "رصيد أول المدة", "DR": "مدين", "CR": "دائن", "PERM_E": "تعديل", "PERM_V": "عرض فقط",
    "MULTI_CELLS": "تعديل عدة خلايا:",
    "ST_DRAFT": "مسودة", "ST_PENDING": "معلق", "ST_APPROVED": "معتمد", "ST_REJECTED": "مرفوض", "ST_RETURNED": "مرتجع",
    "ACT_APPROVE": "اعتماد", "ACT_REJECT": "رفض", "ACT_RETURN": "إرجاع للتعديل",
    "ERR_SESSION": "انتهت الجلسة، سجل دخول تاني.", "ERR_LOGIN": "اسم المستخدم أو كلمة المرور غير صحيحة.",
    "ERR_ROLE": "الدور غير معرف في جدول الأدوار أو الصلاحيات:", "ERR_NO_PERM": "مش مسموح لك بالعملية دي.",
    "ERR_UNEXPECTED": "حصل خطأ غير متوقع:", "ERR_OLD_PASS": "كلمة المرور الحالية غلط.",
    "ERR_PASS_SHORT": "كلمة المرور لازم 4 حروف على الأقل.", "ERR_PASS_MATCH": "كلمة المرور الجديدة والتأكيد مش متطابقين.",
    "ERR_DUP_USER": "اسم المستخدم متكرر:", "ERR_NO_PASS": "مستخدم من غير كلمة مرور:",
    "ERR_NO_MASTER": "لازم يفضل مستخدم ماستر واحد نشط على الأقل.",
    "ERR_PERM_VALUE": "قيمة الصلاحية لازم E أو V أو N في سطر:",
    "ERR_DATE": "اكتب تاريخ صحيح.", "ERR_DESC": "اكتب بيان القيد.",
    "ERR_LINE_ACC": "سطر من غير كود حساب، رقم:", "ERR_LINE_POSTING": "الحساب مش فرعي (لازم حساب يقبل حركة)، سطر:",
    "ERR_LINE_PARTY": "الطرف مش موجود في شيت العملاء والموردين، سطر:",
    "ERR_LINE_AMT": "كل سطر لازم يبقى مدين أو دائن بس وبقيمة موجبة، سطر:",
    "ERR_TWO_LINES": "القيد لازم سطرين على الأقل.", "ERR_UNBALANCED": "القيد مش متزن: إجمالي المدين لا يساوي الدائن.",
    "ERR_NOT_EDITABLE": "القيد ده متعدل عليه أو معتمد ومينفعش يتعدل.", "ERR_NOT_FOUND": "رقم القيد مش موجود.",
    "ERR_TOO_MANY": "القيد فيه سطور أكتر من شاشة الإدخال.", "ERR_SELECT_ROW": "اختار سطر من الجدول الأول.",
    "ERR_HAS_CHILDREN": "الحساب ليه حسابات فرعية ومينفعش يتحذف.", "ERR_ACC_USED": "الحساب عليه قيود ومينفعش يتحذف.",
    "ERR_LEDGER_PARAMS": "اختار حساب أو طرف.", "ERR_NO_ACCOUNT": "كود الحساب مش موجود في الشجرة.",
    "ASK_DELETE": "متأكد من حذف الحساب؟",
    "OK_PASS": "تم تغيير كلمة المرور.", "OK_SAVED": "تم الحفظ.", "OK_PERMS": "تم تطبيق الصلاحيات.",
    "OK_JE_SAVED": "تم حفظ القيد رقم", "OK_NO_PENDING": "مفيش قيود معلقة.", "OK_APPR_DONE": "تم تنفيذ إجراءات على قيود عددها",
    "WARN_APPR_SKIPPED": "سطور اتسابت (ملاحظة ناقصة للرفض/الإرجاع، أو قيدك انت، أو مش معلق):",
    "IS_TITLE": "قائمة الدخل عن الفترة", "IS_SALES": "المبيعات", "IS_NET_SALES": "صافي المبيعات",
    "IS_COGS": "تكلفة المبيعات", "IS_TOTAL_COGS": "إجمالي تكلفة المبيعات", "IS_GROSS": "مجمل الربح",
    "IS_GA": "المصروفات العمومية والإدارية", "IS_TOTAL_GA": "إجمالي العمومية والإدارية",
    "IS_SELL": "مصروفات البيع والتسويق", "IS_TOTAL_SELL": "إجمالي البيع والتسويق", "IS_OPERATING": "ربح التشغيل",
    "IS_OTHER_INC": "إيرادات أخرى", "IS_TOTAL_OTHER_INC": "إجمالي الإيرادات الأخرى",
    "IS_OTHER_EXP": "مصروفات أخرى", "IS_TOTAL_OTHER_EXP": "إجمالي المصروفات الأخرى", "IS_NET": "صافي الربح (الخسارة)",
    "BS_TITLE": "الميزانية (المركز المالي) في", "BS_ASSETS": "الأصول", "BS_TOTAL_ASSETS": "إجمالي الأصول",
    "BS_LIAB": "الخصوم", "BS_TOTAL_LIAB": "إجمالي الخصوم", "BS_EQUITY": "حقوق الملكية",
    "BS_TOTAL_EQUITY": "إجمالي حقوق الملكية", "BS_UNCLOSED_PL": "أرباح (خسائر) لم تقفل بعد",
    "BS_TOTAL_LE": "إجمالي الخصوم وحقوق الملكية", "BS_CHECK": "فرق التوازن (لازم صفر)",
    "LOG_LOGIN": "دخول", "LOG_LOGIN_FAIL": "محاولة دخول فاشلة", "LOG_LOGOUT": "خروج",
    "LOG_PASS_CHANGE": "تغيير كلمة المرور", "LOG_PASS_SET": "تعيين كلمة مرور", "LOG_PERMS_APPLY": "تطبيق الصلاحيات",
    "LOG_ROW_ADD": "إضافة سطر", "LOG_ACC_DELETE": "حذف حساب", "LOG_JE_DRAFT": "حفظ قيد مسودة",
    "LOG_JE_SUBMIT": "إرسال قيد للاعتماد", "LOG_JE_APPROVE": "اعتماد قيد", "LOG_JE_REJECT": "رفض قيد",
    "LOG_JE_RETURN": "إرجاع قيد", "LOG_REPORT": "تحديث تقرير", "LOG_EDIT": "تعديل خلية", "LOG_PDF": "تصدير PDF",
    "LOG_PRINT": "طباعة", "LOG_OPENING": "تحميل الأرصدة الافتتاحية",
    "LOG_DOC_SUBMIT": "إرسال مستند للاعتماد", "LOG_DOC_DRAFT": "حفظ مستند مسودة", "LOG_IMPORT": "نقل حركات من الشيت القديم",
    "ERR_LINKED_DOC": "القيد ده تابع لمستند (فاتورة أو إذن مخزن)، عدّله من شاشة المستند نفسه.",
    "ERR_DOC_TYPE": "اختار نوع المستند.", "ERR_PAY": "اختار طريقة السداد.",
    "ERR_PARTY_REQUIRED": "اختار الطرف (العميل أو المورد) من القائمة.",
    "ERR_ITEM": "كود الصنف مش موجود في شيت الأصناف، سطر:", "ERR_QTY": "الكمية لازم أكبر من صفر، سطر:",
    "ERR_DIR": "في التضريبة حدد لكل سطر وارد أو صادر، سطر:", "ERR_PRICE": "اكتب سعر الشراء، سطر:",
    "ERR_NO_LINES": "مفيش سطور في المستند.", "ERR_CMP_LINES": "التضريبة لازم فيها سطر وارد وسطر صادر على الأقل.",
    "ERR_NO_COST": "مفيش رصيد أو تكلفة للصنف عشان يتصرف بالمتوسط:", "ERR_SAME_ACC": "الحساب المقابل هو نفس حساب النقدية.",
    "OK_DOC_SAVED": "تم حفظ المستند", "OK_NOTHING": "مفيش حاجة تتسجل.",
    "DIR_IN": "وارد", "DIR_OUT": "صادر", "ORDER": "طلبية", "DISPATCH": "إذن", "BALANCE": "الرصيد",
    "SL_INVOICE": "فاتورة بيع", "SL_RETURN": "مرتجع مبيعات", "SL_DISCOUNT": "خصم مسموح به",
    "VC_RECEIPT": "سند قبض", "VC_PAYMENT": "سند صرف",
    "OWN_FACTORY": "ملك المصنع", "OWN_CUSTOMER": "ملك العميل", "MOLD_JE": "استهلاك أسطمبات شهر",
}
DOC_TYPES = [("PUR", "فاتورة شراء", "2101"), ("ISS", "صرف خامات لأوردر (تكلفة مبيعات)", "5101"),
             ("SPR", "صرف قطع غيار ومهمات للتشغيل", "5205"), ("CMP", "تضريبة / تحويل داخلي", ""),
             ("ADJ+", "تسوية جرد بالزيادة", "4303"), ("ADJ-", "تسوية جرد بالعجز", "5403"),
             ("OPN", "رصيد افتتاحي", "3303")]
DOC_LABEL = {k: n for k, n, _ in DOC_TYPES}
PAY = [("آجل", ""), ("الخزينة", "1501"), ("فودافون كاش", "1502"), ("إنستاباي", "1503"), ("البنك", "1504")]

BTN_HOME = "الرئيسية"
BUTTONS = [
    ("WARN", "B9:B10", "فتح شاشة الدخول", "Auto_Open"),
    ("LOGIN", "C10:C11", "دخول", "BtnLogin"),
    ("HOME", "E6:E7", "تسجيل خروج", "BtnLogout"),
    ("HOME", "F15:F16", "تغيير كلمة المرور", "BtnChangePassword"),
    ("JE_ENTRY", "I8:I9", "حفظ مسودة", "BtnJESaveDraft"),
    ("JE_ENTRY", "I11:I12", "إرسال للاعتماد", "BtnJESubmit"),
    ("JE_ENTRY", "I14:I15", "تحميل القيد", "BtnJELoad"),
    ("JE_ENTRY", "I17:I18", "قيد جديد / مسح", "BtnJEClear"),
    ("JE_ENTRY", "I20:I21", BTN_HOME, "GoHome"),
    ("STOCK_ENTRY", "L11:L12", "حفظ مسودة", "BtnStockDraft"),
    ("STOCK_ENTRY", "L14:L15", "إرسال للاعتماد", "BtnStockSubmit"),
    ("STOCK_ENTRY", "L17:L18", "تحميل الإذن", "BtnStockLoad"),
    ("STOCK_ENTRY", "L20:L21", "إذن جديد / مسح", "BtnStockClear"),
    ("STOCK_ENTRY", "L23:L24", BTN_HOME, "GoHome"),
    ("SALES_ENTRY", "J12:J13", "حفظ مسودة", "BtnSaleDraft"),
    ("SALES_ENTRY", "J15:J16", "إرسال للاعتماد", "BtnSaleSubmit"),
    ("SALES_ENTRY", "J18:J19", "تحميل الفاتورة", "BtnSaleLoad"),
    ("SALES_ENTRY", "J21:J22", "فاتورة جديدة / مسح", "BtnSaleClear"),
    ("SALES_ENTRY", "J24:J25", BTN_HOME, "GoHome"),
    ("VOUCHER", "I4:I5", "إرسال للاعتماد", "BtnVoucherSubmit"),
    ("VOUCHER", "I7:I8", BTN_HOME, "GoHome"),
    ("APPROVALS", "R6:R7", "تحديث القائمة", "BtnApprRefresh"),
    ("APPROVALS", "R9:R10", "تنفيذ الإجراءات", "BtnApprExecute"),
    ("APPROVALS", "R12:R13", "عرض سطور القيد المحدد", "BtnApprLines"),
    ("APPROVALS", "R15:R16", BTN_HOME, "GoHome"),
    ("TB", "N4:N5", "تحديث", "BtnTBRefresh"),
    ("LEDGER", "K4:K5", "تحديث", "BtnLedgerRefresh"),
    ("FS", "J4:J5", "تحديث", "BtnFSRefresh"),
    ("ITEM_CARD", "P4:P5", "تحديث", "BtnItemCard"),
    ("PARTY_BAL", "J4:J5", "تحديث", "BtnPartyBalances"),
    ("SALES_REP", "M4:M5", "تحديث", "BtnSalesReport"),
    ("MOLD_USE", "M3:M4", "قيد استهلاك الشهر", "BtnMoldAmortize"),
]
REPORT_BUTTONS = [("JE", "R", 3), ("TB", "N", 7), ("LEDGER", "K", 7), ("FS", "J", 7), ("COA", "L", 3),
                  ("PARTIES", "J", 3), ("ITEMS", "N", 3), ("ASSETS", "M", 3), ("LOG", "J", 3), ("STOCK", "R", 3),
                  ("SALES", "P", 3), ("ITEM_CARD", "P", 7), ("PARTY_BAL", "J", 7), ("SALES_REP", "M", 7),
                  ("MOLDS", "R", 3), ("MOLD_USE", "M", 6)]
ADD_ROW = {"COA", "PARTIES", "ITEMS", "ASSETS", "MOLDS", "MOLD_USE"}
for key, col, start in REPORT_BUTTONS:
    if key in ADD_ROW:
        BUTTONS.append((key, f"{col}{start}:{col}{start + 1}", "إضافة سطر", "BtnAddRow"))
        start += 3
    if key == "COA":
        BUTTONS.append((key, f"{col}{start}:{col}{start + 1}", "ترتيب بالكود", "BtnCOASort"))
        BUTTONS.append((key, f"{col}{start + 3}:{col}{start + 4}", "حذف الحساب المحدد", "BtnCOADelete"))
        start += 6
    BUTTONS.append((key, f"{col}{start}:{col}{start + 1}", "تصدير PDF", "BtnExportPDF"))
    BUTTONS.append((key, f"{col}{start + 3}:{col}{start + 4}", "طباعة", "BtnPrint"))
    BUTTONS.append((key, f"{col}{start + 6}:{col}{start + 7}", BTN_HOME, "GoHome"))
BUTTONS += [
    ("USERS", "J3:J4", "إضافة مستخدم", "BtnAddRow"),
    ("USERS", "J6:J7", "حفظ كلمات المرور", "BtnSavePasswords"),
    ("USERS", "J9:J10", BTN_HOME, "GoHome"),
    ("PERMS", "J4:J5", "تطبيق الصلاحيات", "BtnApplyPerms"),
    ("PERMS", "J7:J8", BTN_HOME, "GoHome"),
    ("SETTINGS", "E3:E4", BTN_HOME, "GoHome"),
]
EDIT_COLS = [("COA", "tblCOA", "1,2,3,9"), ("USERS", "tblUsers", "1,2,3,4,6"), ("PERMS", "tblPerms", "4,5,6"),
             ("PARTIES", "tblParties", "1,2,3,4,7"), ("ITEMS", "tblItems", "1,2,3,4,11"),
             ("ASSETS", "tblAssets", "1,2,4,10"), ("MOLDS", "tblMolds", "1,2,3,4,5,6,7,8,9,15"),
             ("MOLD_USE", "tblMoldUse", "1,2,4,5,6,7,10")]
PRINT_COLS = [("JE", "A:O"), ("TB", "B:J"), ("LEDGER", "B:I"), ("FS", "B:H"), ("COA", "B:J"), ("LOG", "A:H"),
              ("APPROVALS", "B:P"), ("JE_ENTRY", "B:G"), ("PARTIES", "B:H"), ("ITEMS", "B:L"), ("ASSETS", "B:K"),
              ("STOCK", "A:P"), ("SALES", "A:N"), ("ITEM_CARD", "B:N"), ("PARTY_BAL", "B:G"), ("SALES_REP", "B:J"),
              ("MOLDS", "B:P"), ("MOLD_USE", "B:K"), ("STOCK_ENTRY", "B:J"), ("SALES_ENTRY", "B:H"),
              ("VOUCHER", "B:G")]


def pass_hash(user: str, pwd: str) -> str:
    """Same algorithm as PassHash in modNewStar.bas."""
    m, h = 2147483629, 5381
    for ch in user.lower() + "|" + pwd:
        h = (h * 33 + (ord(ch) & 0xFFFF)) % m
    return f"H{h}"


# ---------------------------------------------------------------- styling helpers

def sheet(wb, key, widths=None):
    ws = wb.create_sheet(NAME[key])
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    for col, w in (widths or {}).items():
        ws.column_dimensions[col].width = w
    return ws


def title(ws, cell, text, size=16):
    c = ws[cell]
    c.value = text
    c.font = Font(name=FONT, bold=True, size=size, color=NAVY)


def label(ws, cell, text):
    c = ws[cell]
    c.value = text
    c.font = Font(name=FONT, bold=True)
    c.alignment = Alignment(horizontal="right", vertical="center")


def inp(ws, ref, value=None, fmt=None):
    for row in ws[ref] if ":" in ref else [[ws[ref]]]:
        for c in row:
            c.fill = PatternFill("solid", fgColor=INPUT)
            c.border = BOX
            c.protection = Protection(locked=False)
            c.font = Font(name=FONT)
            if fmt:
                c.number_format = fmt
    if value is not None:
        ws[ref.split(":")[0]].value = value


def header_row(ws, row, col, names):
    for i, n in enumerate(names):
        c = ws.cell(row=row, column=col + i, value=n)
        c.font = Font(name=FONT, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def add_table(ws, name, top_left_row, top_left_col, headers, rows, formulas=None, text_cols=(), formats=None):
    """Excel table; formulas = {col_index(1-based): structured formula without '='}."""
    formulas, formats = formulas or {}, formats or {}
    header_row(ws, top_left_row, top_left_col, headers)
    rows = rows or [[None] * len(headers)]
    for r_i, r in enumerate(rows, start=1):
        for c_i in range(1, len(headers) + 1):
            c = ws.cell(row=top_left_row + r_i, column=top_left_col + c_i - 1)
            if c_i in formulas:
                c.value = "=" + formulas[c_i]
            else:
                v = r[c_i - 1] if c_i - 1 < len(r) else None
                c.value = None if (isinstance(v, float) and pd.isna(v)) else v
            c.font = Font(name=FONT)
            c.border = BOX
            if c_i in text_cols:
                c.number_format = "@"
            if c_i in formats:
                c.number_format = formats[c_i]
    last_col = ws.cell(row=1, column=top_left_col + len(headers) - 1).column_letter
    first_col = ws.cell(row=1, column=top_left_col).column_letter
    ref = f"{first_col}{top_left_row}:{last_col}{top_left_row + len(rows)}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    t.tableColumns = [TableColumn(id=i, name=h, calculatedColumnFormula=TableFormula(attr_text=formulas[i])
                                  if i in formulas else None) for i, h in enumerate(headers, start=1)]
    ws.add_table(t)
    return t


def wb_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def local_name(ws, name, ref):
    ws.defined_names[name] = DefinedName(name, attr_text=ref)


def q(key):
    return "'" + NAME[key] + "'"


def list_validation(ws, ref, formula):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error, dv.errorTitle = "اختار من القائمة", "New Star"
    ws.add_data_validation(dv)
    dv.add(ref)


# ---------------------------------------------------------------- opening data

def load_opening(path):
    x = pd.read_excel(path, sheet_name=None)
    acc, parties, stock, assets = (x["أرصدة الحسابات"], x["العملاء والموردين"], x["المخزون"], x["الأصول الثابتة"])
    return acc, parties, stock, assets


def opening_lines(acc, parties, items, assets):
    """Opening journal lines: (code, party, dr, cr, note). Difference closes on 3303."""
    lines = []
    for _, r in acc.iterrows():
        lines.append((str(int(r["كود الحساب"])), "", r["مدين"] if pd.notna(r["مدين"]) else 0,
                      r["دائن"] if pd.notna(r["دائن"]) else 0, r["البيان"]))
    for _, r in parties.iterrows():
        dr = r["مدين"] if pd.notna(r["مدين"]) else 0
        cr = r["دائن"] if pd.notna(r["دائن"]) else 0
        if dr:
            lines.append(("2101" if r["النوع"] == "مورد" else "1401", r["الاسم"], dr, 0, "رصيد افتتاحي"))
        if cr:
            lines.append(("1401" if r["النوع"] == "عميل" else "2101", r["الاسم"], 0, cr, "رصيد افتتاحي"))
    by_acc = {}
    for it in items:
        by_acc[it["account"]] = by_acc.get(it["account"], 0) + round(it["qty"] * it["cost"], 2)
    for code, v in sorted(by_acc.items()):
        lines.append((code, "", round(v, 2), 0, "مخزون أول المدة حسب شيت الأصناف"))
    for code, v in sorted(assets.groupby("كود حساب الفئة")["التكلفة"].sum().items()):
        lines.append((str(code), "", float(v), 0, "أصول ثابتة حسب شيت الأصول"))
    diff = round(sum(l[2] for l in lines) - sum(l[3] for l in lines), 2)
    if diff:
        lines.append(("3303", "", -diff if diff < 0 else 0, diff if diff > 0 else 0, "فرق الأرصدة الافتتاحية"))
    return lines


# ---------------------------------------------------------------- build

def build(opening_path, wb_pickle, out_path, notes_path=None):
    acc, parties, stock, assets = load_opening(opening_path)
    group_acc = {g: a for g, a, _ in GROUPS}
    group_pfx = {g: p for g, _, p in GROUPS}
    counters, items = {}, []
    for _, r in stock.iterrows():
        g = r["المجموعة"]
        counters[g] = counters.get(g, 0) + 1
        items.append({"code": f"{group_pfx[g]}-{counters[g]:04d}", "name": r["اسم الصنف"], "group": g,
                      "unit": r["الوحدة"], "account": group_acc[g], "qty": float(r["الكمية"]),
                      "cost": float(r["تكلفة الوحدة"])})
    opening_items = [dict(i) for i in items]

    # ---- April-July movements from the original workbook
    hist = History(wb_pickle, items, list(parties["الاسم"]))
    docs = hist.build(group_acc)
    for it in hist.items.values():
        if it.get("new"):
            g = it["group"]
            counters[g] = counters.get(g, 0) + 1
            it["code"] = f"{group_pfx[g]}-{counters[g]:04d}"
    all_items = list(hist.items.values())
    stamp = datetime(2026, 10, 15, 12, 0)
    names = {a[0]: a[1] for a in ACCOUNTS}
    lines = opening_lines(acc, parties, opening_items, assets)
    approved = MSG["ST_APPROVED"]
    imp_note = "منقول من شيت zyad_15-07-2026"

    def je_row(no, dt, desc, l, by="master", note="مرحّل من ملف الأرصدة الافتتاحية"):
        c, p, dr, cr, n = l
        return [no, datetime.combine(dt, datetime.min.time()), desc, c, names[c], p or None, dr or None, cr or None,
                n or None, approved, by, stamp, by, stamp, note]

    je_rows = [je_row("JE-000001", OPENING_DATE, "قيد الأرصدة الافتتاحية 31-03-2026", l) for l in lines]
    st_rows = [["ST-000001", datetime.combine(OPENING_DATE, datetime.min.time()), DOC_LABEL["OPN"], i["code"], i["name"],
                i["qty"], None, i["cost"], round(i["qty"] * i["cost"], 2), None, None, None, "رصيد أول المدة",
                "JE-000001", approved, "master"] for i in opening_items if i["qty"] > 0]
    sl_rows, new_assets = [], []
    n_je, n_st, n_si = 1, 1, 0
    kind_type = {"purchase": "PUR", "issue": "ISS", "compound": "CMP"}
    for d in docs:
        n_je += 1
        je = f"JE-{n_je:06d}"
        desc = d["desc"]
        if d.get("stock"):
            n_st += 1
            st = f"ST-{n_st:06d}"
            k = kind_type.get(d["kind"])
            desc = f"{DOC_LABEL[k]} {st} - {d['desc']}" if k else f"{st} - {d['desc']}"
            for tk, it, qi, qo, cost, vi, vo, party, order, note in d["stock"]:
                st_rows.append([st, datetime.combine(d["date"], datetime.min.time()), DOC_LABEL[tk], it["code"],
                                it["name"], qi or None, qo or None, cost, vi or None, vo or None, party or None,
                                order or None, note or d["desc"], je, approved, "import"])
        if d.get("sales"):
            n_si += 1
            si = f"SI-{n_si:06d}"
            ret = d["kind"] == "return"
            lab = MSG["SL_RETURN"] if ret else MSG["SL_INVOICE"]
            desc = f"{lab} {si} - {MSG['DISPATCH']} {d['dispatch']} - {d['party']}"
            for ds, order, unit, qty, price, val in d["sales"]:
                sg = -1 if ret else 1
                sl_rows.append([si, datetime.combine(d["date"], datetime.min.time()), lab, d["party"], d["dispatch"],
                                order or None, ds, unit, sg * qty, price, sg * val, je, approved, "import"])
        for nm, code, v in d.get("assets", []):
            new_assets.append((nm, code, d["date"], v))
        je_rows += [je_row(je, d["date"], desc, l, "import", imp_note) for l in d["lines"]]
    molds, usage = hist.molds_and_usage()

    wb = Workbook()
    wb.remove(wb.active)

    # ---- WARN (visible on a file opened without macros)
    ws = sheet(wb, "WARN", {"B": 26})
    ws.row_dimensions[9].height = ws.row_dimensions[10].height = 22
    title(ws, "B2", "نظام New Star لإدارة الحسابات", 18)
    for i, t in enumerate([
        "النظام ده شغال بالماكرو. لو انت شايف الشيت ده بس، يبقى الماكرو مقفول.",
        "اضغط Enable Content في الشريط الأصفر فوق، أو اقفل الملف وافتحه تاني بعد تفعيل الماكرو.",
        "لو الملف جالك من الإنترنت أو الإيميل: كليك يمين على الملف > Properties > علّم Unblock > OK.",
        "لو فعّلت الماكرو وشاشة الدخول مظهرتش، اضغط الزر اللي تحت أو اقفل الملف وافتحه تاني.",
    ], start=4):
        ws[f"B{i}"].value = t
        ws[f"B{i}"].font = Font(name=FONT, size=13)

    # ---- LOGIN
    ws = sheet(wb, "LOGIN", {"A": 4, "B": 20, "C": 32})
    ws["B2"].value = "=setCompany"
    ws["B2"].font = Font(name=FONT, bold=True, size=20, color=NAVY)
    title(ws, "B4", "تسجيل الدخول", 14)
    label(ws, "B6", "المستخدم")
    inp(ws, "C6", "master")
    label(ws, "B8", "كلمة المرور")
    inp(ws, "C8", None, ";;;")
    for r in (6, 8, 10, 11):
        ws.row_dimensions[r].height = 24
    ws["B13"].value = "كلمة المرور مش بتظهر في الخلية. بعد الدخول بتتمسح تلقائي."
    ws["B13"].font = Font(name=FONT, italic=True, color="808080")
    list_validation(ws, "C6", 'INDIRECT("tblUsers[المستخدم]")')
    wb_name(wb, "lgnUser", f"{q('LOGIN')}!$C$6")
    wb_name(wb, "lgnPass", f"{q('LOGIN')}!$C$8")
    local_name(ws, "InputArea", f"{q('LOGIN')}!$C$6,{q('LOGIN')}!$C$8")
    local_name(ws, "NoLog", f"{q('LOGIN')}!$C$6:$C$8")

    # ---- HOME
    ws = sheet(wb, "HOME", {"A": 4, "B": 30, "C": 14, "D": 4, "E": 22, "F": 22})
    ws["B2"].value = "=setCompany"
    ws["B2"].font = Font(name=FONT, bold=True, size=20, color=NAVY)
    ws["B3"].font = Font(name=FONT, bold=True, size=12)
    title(ws, "B5", "الشاشات المتاحة لك", 12)
    title(ws, "E10", "تغيير كلمة المرور", 12)
    for r, t in ((11, "الحالية"), (12, "الجديدة"), (13, "تأكيد الجديدة")):
        label(ws, f"E{r}", t)
        inp(ws, f"F{r}", None, ";;;")
    wb_name(wb, "hmWelcome", f"{q('HOME')}!$B$3")
    wb_name(wb, "hmOld", f"{q('HOME')}!$F$11")
    wb_name(wb, "hmNew", f"{q('HOME')}!$F$12")
    wb_name(wb, "hmNew2", f"{q('HOME')}!$F$13")
    local_name(ws, "InputArea", f"{q('HOME')}!$F$11:$F$13")
    local_name(ws, "NoLog", f"{q('HOME')}!$F$11:$F$13")

    # ---- JE_ENTRY
    ws = sheet(wb, "JE_ENTRY", {"A": 3, "B": 13, "C": 30, "D": 24, "E": 15, "F": 15, "G": 30, "H": 3, "I": 24})
    title(ws, "B1", "قيد يومية جديد")
    label(ws, "B3", "تحميل قيد رقم")
    inp(ws, "C3")
    label(ws, "E3", "الحالة")
    ws["F3"].font = Font(name=FONT, bold=True, color="C00000")
    label(ws, "B4", "رقم القيد")
    ws["C4"].font = Font(name=FONT, bold=True)
    label(ws, "E4", "التاريخ")
    inp(ws, "F4", datetime.combine(date.today(), datetime.min.time()), "yyyy-mm-dd")
    label(ws, "B5", "البيان")
    ws.merge_cells("C5:G5")
    inp(ws, "C5:G5")
    label(ws, "B6", "ملاحظة المعتمد")
    ws.merge_cells("C6:G6")
    ws["C6"].font = Font(name=FONT, color="C00000")
    header_row(ws, 7, 2, ["كود الحساب", "اسم الحساب", "الطرف (عميل/مورد)", "مدين", "دائن", "بيان السطر"])
    for r in range(8, 28):
        inp(ws, f"B{r}", None, "@")
        ws[f"C{r}"].value = f'=IF(B{r}="","",IFERROR(INDEX(tblCOA[اسم الحساب],MATCH(B{r},tblCOA[الكود],0)),"؟ كود غير موجود"))'
        ws[f"C{r}"].border = BOX
        inp(ws, f"D{r}")
        inp(ws, f"E{r}", None, MONEY)
        inp(ws, f"F{r}", None, MONEY)
        inp(ws, f"G{r}")
    label(ws, "C28", "الإجمالي")
    ws["E28"].value, ws["F28"].value = "=SUM(E8:E27)", "=SUM(F8:F27)"
    label(ws, "C29", "الفرق")
    ws["E29"].value = "=ROUND(E28-F28,2)"
    for c in ("E28", "F28", "E29"):
        ws[c].number_format = MONEY
        ws[c].font = Font(name=FONT, bold=True)
    list_validation(ws, "B8:B27", 'INDIRECT("tblCOA[الكود]")')
    list_validation(ws, "D8:D27", 'INDIRECT("tblParties[الاسم]")')
    for k, ref in (("jeeLoad", "$C$3"), ("jeeStatus", "$F$3"), ("jeeNo", "$C$4"), ("jeeDate", "$F$4"),
                   ("jeeDesc", "$C$5"), ("jeeNote", "$C$6"), ("jeeLines", "$B$8:$G$27")):
        wb_name(wb, k, f"{q('JE_ENTRY')}!{ref}")
    local_name(ws, "InputArea", ",".join(f"{q('JE_ENTRY')}!{r}" for r in
                                         ("$C$3", "$F$4", "$C$5:$G$5", "$B$8:$B$27", "$D$8:$G$27")))

    # ---- JE register with the opening journal
    ws = sheet(wb, "JE", {"A": 12, "B": 11, "C": 28, "D": 9, "E": 26, "F": 22, "G": 14, "H": 14, "I": 26,
                          "J": 9, "K": 9, "L": 16, "M": 9, "N": 16, "O": 20, "P": 3, "Q": 3, "R": 22})
    title(ws, "A1", "دفتر اليومية (كل السطور)")
    add_table(ws, "tblJE", 3, 1, ["رقم القيد", "التاريخ", "البيان العام", "كود الحساب", "اسم الحساب", "الطرف",
                                  "مدين", "دائن", "بيان السطر", "الحالة", "أنشأه", "تاريخ الإنشاء", "اعتمده",
                                  "تاريخ الاعتماد", "ملاحظة الاعتماد"],
              je_rows, text_cols=(4,), formats={2: "yyyy-mm-dd", 7: MONEY, 8: MONEY, 12: "yyyy-mm-dd hh:mm",
                                                14: "yyyy-mm-dd hh:mm"})
    ws.freeze_panes = "A4"
    for nm, col in (("jeDate", "B"), ("jeCode", "D"), ("jeDr", "G"), ("jeCr", "H"), ("jeStatus", "J"),
                    ("jeParty", "F"), ("jeNote", "I")):
        wb_name(wb, nm, f"{q('JE')}!${col}:${col}")

    # ---- APPROVALS
    ws = sheet(wb, "APPROVALS", {"A": 3, "B": 12, "C": 11, "D": 30, "E": 14, "F": 10, "G": 16, "H": 16, "I": 28,
                                 "J": 3, "K": 12, "L": 9, "M": 24, "N": 13, "O": 13, "P": 26, "Q": 3, "R": 26})
    title(ws, "B1", "اعتماد القيود")
    ws["B3"].value = ("اضغط «تحديث القائمة»، اختار الإجراء لكل قيد (اعتماد / رفض / إرجاع للتعديل)، "
                      "والملاحظة لازمة مع الرفض والإرجاع، وبعدين «تنفيذ الإجراءات».")
    ws["B3"].font = Font(name=FONT, italic=True, color="595959")
    header_row(ws, 5, 2, ["رقم القيد", "التاريخ", "البيان", "الإجمالي", "أنشأه", "تاريخ الإنشاء", "الإجراء", "الملاحظة"])
    header_row(ws, 5, 11, ["رقم القيد", "الكود", "الحساب", "مدين", "دائن", "الطرف / البيان"])
    for r in range(6, 206):
        ws[f"C{r}"].number_format = "yyyy-mm-dd"
        ws[f"E{r}"].number_format = MONEY
        ws[f"G{r}"].number_format = "yyyy-mm-dd hh:mm"
        ws[f"N{r}"].number_format = MONEY
        ws[f"O{r}"].number_format = MONEY
    inp(ws, "H6:I205")
    list_validation(ws, "H6:H205", "lstApprActions")
    local_name(ws, "InputArea", f"{q('APPROVALS')}!$H$6:$I$205")
    ws.freeze_panes = "A6"

    # ---- TB
    ws = sheet(wb, "TB", {"A": 3, "B": 10, "C": 36, "D": 7, "E": 15, "F": 15, "G": 15, "H": 15, "I": 15, "J": 15,
                          "M": 3, "N": 18})
    title(ws, "B1", "ميزان المراجعة")
    label(ws, "B4", "من")
    inp(ws, "C4", datetime(2026, 4, 1), "yyyy-mm-dd")
    label(ws, "D4", "إلى")
    inp(ws, "E4", datetime(2026, 12, 31), "yyyy-mm-dd")
    label(ws, "F4", "المستوى")
    inp(ws, "G4", 4)
    list_validation(ws, "G4", '"1,2,3,4"')
    header_row(ws, 6, 5, ["رصيد أول المدة", "", "حركة الفترة", "", "رصيد آخر المدة", ""])
    for rng in ("E6:F6", "G6:H6", "I6:J6"):
        ws.merge_cells(rng)
    header_row(ws, 7, 2, ["الكود", "اسم الحساب", "مستوى", "مدين", "دائن", "مدين", "دائن", "مدين", "دائن"])
    ws.column_dimensions["K"].hidden = ws.column_dimensions["L"].hidden = True
    for k, ref in (("tbFrom", "$C$4"), ("tbTo", "$E$4"), ("tbLevel", "$G$4")):
        wb_name(wb, k, f"{q('TB')}!{ref}")
    local_name(ws, "ParamArea", f"{q('TB')}!$C$4,{q('TB')}!$E$4,{q('TB')}!$G$4")
    ws.freeze_panes = "A8"

    # ---- LEDGER
    ws = sheet(wb, "LEDGER", {"A": 3, "B": 12, "C": 14, "D": 9, "E": 46, "F": 15, "G": 15, "H": 15, "I": 9,
                              "J": 3, "K": 18})
    title(ws, "B1", "كشف حساب / كشف عميل أو مورد")
    label(ws, "B4", "الحساب")
    inp(ws, "C4", "1501", "@")
    label(ws, "D4", "من")
    inp(ws, "E4", datetime(2026, 4, 1), "yyyy-mm-dd")
    label(ws, "F4", "إلى")
    inp(ws, "G4", datetime(2026, 12, 31), "yyyy-mm-dd")
    label(ws, "B5", "الطرف")
    inp(ws, "C5")
    ws["E5"].value = '=IF(C4="","",IFERROR(INDEX(tblCOA[اسم الحساب],MATCH(C4,tblCOA[الكود],0)),"؟"))'
    ws["E5"].font = Font(name=FONT, bold=True, color=NAVY)
    list_validation(ws, "C4", 'INDIRECT("tblCOA[الكود]")')
    list_validation(ws, "C5", 'INDIRECT("tblParties[الاسم]")')
    header_row(ws, 7, 2, ["التاريخ", "رقم القيد", "الكود", "البيان", "مدين", "دائن", "الرصيد", "طبيعته"])
    for k, ref in (("lgAcc", "$C$4"), ("lgFrom", "$E$4"), ("lgTo", "$G$4"), ("lgParty", "$C$5")):
        wb_name(wb, k, f"{q('LEDGER')}!{ref}")
    local_name(ws, "ParamArea", ",".join(f"{q('LEDGER')}!{r}" for r in ("$C$4", "$E$4", "$G$4", "$C$5")))
    ws.freeze_panes = "A8"

    # ---- FS
    ws = sheet(wb, "FS", {"A": 3, "B": 36, "C": 7, "D": 17, "E": 4, "F": 36, "G": 7, "H": 17, "I": 3, "J": 18})
    title(ws, "B1", "القوائم المالية")
    label(ws, "B4", "من")
    inp(ws, "C4", datetime(2026, 4, 1), "yyyy-mm-dd")
    ws.column_dimensions["C"].width = 12
    label(ws, "D4", "إلى")
    inp(ws, "E4", datetime(2026, 12, 31), "yyyy-mm-dd")
    ws.column_dimensions["E"].width = 12
    wb_name(wb, "fsFrom", f"{q('FS')}!$C$4")
    wb_name(wb, "fsTo", f"{q('FS')}!$E$4")
    local_name(ws, "ParamArea", f"{q('FS')}!$C$4,{q('FS')}!$E$4")

    # ---- COA
    ws = sheet(wb, "COA", {"A": 3, "B": 10, "C": 34, "D": 30, "E": 11, "F": 8, "G": 12, "H": 11, "I": 10, "J": 24,
                           "K": 3, "L": 20})
    title(ws, "B1", "شجرة الحسابات")
    code = "tblCOA[[#This Row],[الكود]]"
    add_table(ws, "tblCOA", 3, 2,
              ["الكود", "اسم الحساب", "الاسم بالإنجليزية", "الحساب الأب", "المستوى", "النوع", "القائمة", "يقبل حركة", "ملاحظات"],
              [[c, ar, en, None, None, None, None, None, None] for c, ar, en, _k, _contra in ACCOUNTS],
              formulas={
                  4: f'IF(LEN({code})<=1,"",LEFT({code},IF(LEN({code})=2,1,IF(LEN({code})<=4,2,4))))',
                  5: f"IF(LEN({code})=1,1,IF(LEN({code})=2,2,IF(LEN({code})<=4,3,4)))",
                  6: f'IFERROR(CHOOSE(VALUE(LEFT({code},1)),"أصول","خصوم","حقوق ملكية","إيرادات","تكاليف","مصروفات","مصروفات","مصروفات"),"")',
                  7: f'IF({code}="","",IF(VALUE(LEFT({code},1))<=3,"الميزانية","قائمة الدخل"))',
                  8: f'IF(COUNTIF(tblCOA[الكود],{code}&"?*")=0,"نعم","لا")',
              }, text_cols=(1,))
    ws.freeze_panes = "A4"

    # ---- PARTIES
    ws = sheet(wb, "PARTIES", {"A": 3, "B": 28, "C": 12, "D": 30, "E": 22, "F": 14, "G": 14, "H": 24, "I": 3, "J": 20})
    title(ws, "B1", "العملاء والموردين")
    add_table(ws, "tblParties", 3, 2, ["الاسم", "النوع", "التليفون", "العنوان", "رصيد افتتاحي مدين",
                                       "رصيد افتتاحي دائن", "ملاحظات"],
              [[r["الاسم"], r["النوع"], r["التليفون"], r["العنوان"], r["مدين"], r["دائن"], None]
               for _, r in parties.iterrows()] +
              [[n, "عميل", None, None, None, None, "اتضاف من حركات المخازن"] for n in hist.new_parties],
              text_cols=(3,), formats={5: MONEY, 6: MONEY})
    list_validation(ws, "C4:C500", '"عميل,مورد,عميل ومورد"')
    ws.freeze_panes = "A4"

    # ---- ITEMS (balances live from the stock register)
    ws = sheet(wb, "ITEMS", {"A": 3, "B": 10, "C": 32, "D": 18, "E": 8, "F": 10, "G": 12, "H": 12, "I": 12, "J": 14,
                             "K": 13, "L": 22, "M": 3, "N": 20})
    title(ws, "B1", "الأصناف وأرصدتها (بالمتوسط المرجح)")
    ws["B2"].value = "الأرصدة بتتحسب من حركة المخزن المعتمدة. أي إضافة أو صرف بيتعمل من «إذن مخزن»."
    ws["B2"].font = Font(name=FONT, italic=True, color="595959")
    c = "tblItems[[#This Row],[كود الصنف]]"
    crit = f"stCode,{c},stStatus,stApproved"
    add_table(ws, "tblItems", 3, 2, ["كود الصنف", "اسم الصنف", "المجموعة", "الوحدة", "حساب المخزون",
                                     "الكمية الواردة", "الكمية المنصرفة", "رصيد الكمية", "رصيد القيمة",
                                     "متوسط التكلفة", "ملاحظات"],
              [[i["code"], i["name"], i["group"], i["unit"]] + [None] * 7 for i in all_items],
              formulas={5: 'IFERROR(INDEX(tblGroups[حساب المخزون],MATCH(tblItems[[#This Row],[المجموعة]],tblGroups[المجموعة],0)),"")',
                        6: f"SUMIFS(stInQ,{crit})",
                        7: f"SUMIFS(stOutQ,{crit})",
                        8: "ROUND(tblItems[[#This Row],[الكمية الواردة]]-tblItems[[#This Row],[الكمية المنصرفة]],4)",
                        9: f"ROUND(SUMIFS(stInV,{crit})-SUMIFS(stOutV,{crit}),2)",
                        10: "IF(tblItems[[#This Row],[رصيد الكمية]]>0.0001,tblItems[[#This Row],[رصيد القيمة]]/tblItems[[#This Row],[رصيد الكمية]],0)"},
              text_cols=(1,), formats={6: "#,##0.###", 7: "#,##0.###", 8: "#,##0.###;[Red]-#,##0.###", 9: MONEY,
                                       10: "#,##0.00##"})
    list_validation(ws, "D4:D2000", 'INDIRECT("tblGroups[المجموعة]")')
    ws.freeze_panes = "A4"

    # ---- ASSETS
    ws = sheet(wb, "ASSETS", {"A": 3, "B": 10, "C": 34, "D": 11, "E": 10, "F": 12, "G": 14, "H": 10, "I": 10,
                              "J": 14, "K": 24, "L": 20})
    title(ws, "B1", "سجل الأصول الثابتة")
    add_table(ws, "tblAssets", 3, 2, ["كود الأصل", "اسم الأصل", "حساب الفئة", "القسم", "تاريخ الشراء", "التكلفة",
                                      "الخردة", "العمر بالشهور", "مجمع الإهلاك الافتتاحي", "ملاحظات"],
              [[f"FA-{i:04d}", r["اسم الأصل"], str(r["كود حساب الفئة"]), r["القسم"], r["تاريخ الشراء"], r["التكلفة"],
                r["الخردة"], r["العمر بالشهور"], r["مجمع الإهلاك"], None]
               for i, (_, r) in enumerate(assets.iterrows(), start=1)] +
              [[f"FA-{len(assets) + i:04d}", nm, code, "المصنع", datetime.combine(dt, datetime.min.time()), v, 0, 60, 0,
                "مشتريات أبريل - يوليو"] for i, (nm, code, dt, v) in enumerate(new_assets, start=1)],
              text_cols=(1, 3), formats={5: "yyyy-mm-dd", 6: MONEY, 7: MONEY, 9: MONEY})
    ws.freeze_panes = "A4"

    # ---- LOG
    ws = sheet(wb, "LOG", {"A": 18, "B": 10, "C": 12, "D": 22, "E": 18, "F": 10, "G": 30, "H": 40, "I": 3, "J": 18})
    title(ws, "A1", "سجل الحركات")
    add_table(ws, "tblLog", 3, 1, ["التاريخ والوقت", "المستخدم", "الدور", "الإجراء", "الشيت", "الخلية",
                                   "القيمة السابقة", "القيمة الجديدة / التفاصيل"],
              [[stamp, "master", "ماستر", MSG["LOG_OPENING"], NAME["JE"], "", "", f"JE-000001 / {len(lines)} سطر"],
               [stamp, "master", "ماستر", MSG["LOG_IMPORT"], NAME["JE"], "", "",
                f"{len(docs)} مستند: JE-000002 لحد JE-{n_je:06d}، {len(st_rows)} حركة مخزن، {len(sl_rows)} سطر مبيعات"]],
              formats={1: "yyyy-mm-dd hh:mm:ss"}, text_cols=(7, 8))
    ws.freeze_panes = "A4"

    # ---- STOCK_ENTRY (purchase, issue, spare parts, compound, count adjustments)
    ws = sheet(wb, "STOCK_ENTRY", {"A": 3, "B": 13, "C": 30, "D": 9, "E": 13, "F": 13, "G": 12, "H": 13, "I": 10,
                                   "J": 26, "K": 3, "L": 24})
    title(ws, "B1", "إذن مخزن (شراء / صرف / تضريبة / تسوية جرد)")
    label(ws, "B3", "تحميل إذن رقم")
    inp(ws, "C3")
    label(ws, "E3", "الحالة")
    ws["F3"].font = Font(name=FONT, bold=True, color="C00000")
    label(ws, "B4", "رقم الإذن")
    ws["C4"].font = Font(name=FONT, bold=True)
    label(ws, "E4", "التاريخ")
    inp(ws, "F4", datetime.combine(date.today(), datetime.min.time()), "yyyy-mm-dd")
    label(ws, "B5", "نوع المستند")
    inp(ws, "C5", DOC_LABEL["ISS"])
    label(ws, "E5", "السداد (للشراء)")
    inp(ws, "F5", PAY[0][0])
    ws.merge_cells("F5:G5")
    label(ws, "B6", "الطرف")
    inp(ws, "C6")
    label(ws, "E6", "رقم الطلبية")
    inp(ws, "F6")
    label(ws, "B7", "البيان")
    ws.merge_cells("C7:J7")
    inp(ws, "C7:J7")
    label(ws, "B8", "ملاحظة المعتمد")
    ws.merge_cells("C8:J8")
    ws["C8"].font = Font(name=FONT, color="C00000")
    header_row(ws, 10, 2, ["كود الصنف", "اسم الصنف", "الوحدة", "الرصيد الحالي", "متوسط التكلفة", "الكمية",
                           "سعر الوحدة (شراء/زيادة)", "وارد/صادر (تضريبة)", "بيان السطر"])
    for r in range(11, 31):
        inp(ws, f"B{r}", None, "@")
        for col, idx, fmt in (("C", "اسم الصنف", None), ("D", "الوحدة", None), ("E", "رصيد الكمية", "#,##0.###"),
                              ("F", "متوسط التكلفة", "#,##0.00##")):
            ws[f"{col}{r}"].value = (f'=IF(B{r}="","",IFERROR(INDEX(tblItems[{idx}],MATCH(B{r},tblItems[كود الصنف],0)),'
                                     f'"{"؟ كود غير موجود" if col == "C" else ""}"))')
            ws[f"{col}{r}"].border = BOX
            if fmt:
                ws[f"{col}{r}"].number_format = fmt
        inp(ws, f"G{r}", None, "#,##0.###")
        inp(ws, f"H{r}", None, "#,##0.00##")
        inp(ws, f"I{r}")
        inp(ws, f"J{r}")
    list_validation(ws, "B11:B30", 'INDIRECT("tblItems[كود الصنف]")')
    list_validation(ws, "C5", "lstDocTypes")
    list_validation(ws, "F5", "lstPay")
    list_validation(ws, "C6", 'INDIRECT("tblParties[الاسم]")')
    list_validation(ws, "I11:I30", "lstDir")
    ws["B32"].value = ("الصرف والعجز بيتحسبوا بالمتوسط المرجح وقت الاعتماد. في الشراء اكتب سعر الوحدة. "
                       "في التضريبة: الصادر بيتصرف بالمتوسط وتكلفته بتتحمل على الوارد.")
    ws["B32"].font = Font(name=FONT, italic=True, color="595959")
    for k, ref in (("stfLoad", "$C$3"), ("stfStatus", "$F$3"), ("stfNo", "$C$4"), ("stfDate", "$F$4"),
                   ("stfType", "$C$5"), ("stfPay", "$F$5"), ("stfParty", "$C$6"), ("stfOrder", "$F$6"),
                   ("stfDesc", "$C$7"), ("stfNote", "$C$8"), ("stfLines", "$B$11:$J$30")):
        wb_name(wb, k, f"{q('STOCK_ENTRY')}!{ref}")
    local_name(ws, "InputArea", ",".join(f"{q('STOCK_ENTRY')}!{r}" for r in
                                         ("$C$3", "$F$4", "$C$5", "$F$5:$G$5", "$C$6", "$F$6", "$C$7:$J$7",
                                          "$B$11:$B$30", "$G$11:$J$30")))

    # ---- STOCK register
    ws = sheet(wb, "STOCK", {"A": 11, "B": 11, "C": 22, "D": 10, "E": 28, "F": 11, "G": 11, "H": 11, "I": 13,
                             "J": 13, "K": 20, "L": 10, "M": 28, "N": 11, "O": 9, "P": 9, "Q": 3, "R": 22})
    title(ws, "A1", "حركة المخزن (كل الأذون)")
    add_table(ws, "tblStock", 3, 1, ["رقم الإذن", "التاريخ", "النوع", "كود الصنف", "اسم الصنف", "كمية واردة",
                                     "كمية منصرفة", "تكلفة الوحدة", "قيمة الوارد", "قيمة المنصرف", "الطرف",
                                     "الطلبية", "البيان", "رقم القيد", "الحالة", "أنشأه"],
              st_rows, text_cols=(4,), formats={2: "yyyy-mm-dd", 6: "#,##0.###", 7: "#,##0.###", 8: "#,##0.00##",
                                                 9: MONEY, 10: MONEY})
    ws.freeze_panes = "A4"
    for nm, col in (("stCode", "D"), ("stInQ", "F"), ("stOutQ", "G"), ("stInV", "I"), ("stOutV", "J"),
                    ("stStatus", "O"), ("stDate", "B")):
        wb_name(wb, nm, f"{q('STOCK')}!${col}:${col}")

    # ---- SALES_ENTRY
    ws = sheet(wb, "SALES_ENTRY", {"A": 3, "B": 34, "C": 12, "D": 9, "E": 11, "F": 12, "G": 14, "H": 24, "I": 3,
                                   "J": 24})
    title(ws, "B1", "فاتورة بيع / مرتجع مبيعات")
    label(ws, "B3", "تحميل فاتورة رقم")
    inp(ws, "C3")
    label(ws, "E3", "الحالة")
    ws["F3"].font = Font(name=FONT, bold=True, color="C00000")
    label(ws, "B4", "رقم الفاتورة")
    ws["C4"].font = Font(name=FONT, bold=True)
    label(ws, "E4", "التاريخ")
    inp(ws, "F4", datetime.combine(date.today(), datetime.min.time()), "yyyy-mm-dd")
    label(ws, "B5", "النوع")
    inp(ws, "C5", MSG["SL_INVOICE"])
    ws.merge_cells("C5:D5")
    label(ws, "E5", "السداد")
    inp(ws, "F5", PAY[0][0])
    ws.merge_cells("F5:G5")
    label(ws, "B6", "العميل")
    inp(ws, "C6")
    ws.merge_cells("C6:D6")
    label(ws, "E6", "رقم الإذن")
    inp(ws, "F6")
    label(ws, "B7", "خصم مسموح به")
    inp(ws, "C7", None, MONEY)
    label(ws, "B8", "البيان")
    ws.merge_cells("C8:H8")
    inp(ws, "C8:H8")
    label(ws, "B9", "ملاحظة المعتمد")
    ws.merge_cells("C9:H9")
    ws["C9"].font = Font(name=FONT, color="C00000")
    header_row(ws, 11, 2, ["الصنف / الوصف", "الطلبية", "الوحدة", "الكمية", "السعر", "القيمة", "ملاحظة"])
    for r in range(12, 32):
        inp(ws, f"B{r}")
        inp(ws, f"C{r}")
        inp(ws, f"D{r}")
        inp(ws, f"E{r}", None, "#,##0.###")
        inp(ws, f"F{r}", None, "#,##0.00##")
        ws[f"G{r}"].value = f'=IF(E{r}="","",ROUND(E{r}*F{r},2))'
        ws[f"G{r}"].number_format = MONEY
        ws[f"G{r}"].border = BOX
        inp(ws, f"H{r}")
    for r, t, f in ((32, "الإجمالي", "=SUM(G12:G31)"), (33, "الخصم", "=N(C7)"), (34, "الصافي", "=G32-G33")):
        label(ws, f"F{r}", t)
        ws[f"G{r}"].value = f
        ws[f"G{r}"].number_format = MONEY
        ws[f"G{r}"].font = Font(name=FONT, bold=True)
    list_validation(ws, "C5", "lstSaleTypes")
    list_validation(ws, "F5", "lstPay")
    list_validation(ws, "C6", 'INDIRECT("tblParties[الاسم]")')
    list_validation(ws, "D12:D31", "lstCounter")
    for k, ref in (("sfLoad", "$C$3"), ("sfStatus", "$F$3"), ("sfNo", "$C$4"), ("sfDate", "$F$4"),
                   ("sfType", "$C$5"), ("sfPay", "$F$5"), ("sfCust", "$C$6"), ("sfDispatch", "$F$6"),
                   ("sfDisc", "$C$7"), ("sfDesc", "$C$8"), ("sfNote", "$C$9"), ("sfLines", "$B$12:$H$31")):
        wb_name(wb, k, f"{q('SALES_ENTRY')}!{ref}")
    local_name(ws, "InputArea", ",".join(f"{q('SALES_ENTRY')}!{r}" for r in
                                         ("$C$3", "$F$4", "$C$5:$D$5", "$F$5:$G$5", "$C$6:$D$6", "$F$6", "$C$7",
                                          "$C$8:$H$8", "$B$12:$F$31", "$H$12:$H$31")))

    # ---- SALES register
    ws = sheet(wb, "SALES", {"A": 11, "B": 11, "C": 14, "D": 24, "E": 9, "F": 10, "G": 30, "H": 8, "I": 10,
                             "J": 11, "K": 13, "L": 11, "M": 9, "N": 9, "O": 3, "P": 22})
    title(ws, "A1", "سجل المبيعات (الفواتير والمرتجعات)")
    add_table(ws, "tblSales", 3, 1, ["رقم الفاتورة", "التاريخ", "النوع", "العميل", "رقم الإذن", "الطلبية",
                                     "الصنف / الوصف", "الوحدة", "الكمية", "السعر", "القيمة", "رقم القيد",
                                     "الحالة", "أنشأه"],
              sl_rows, text_cols=(5,), formats={2: "yyyy-mm-dd", 9: "#,##0.###", 10: "#,##0.00##", 11: MONEY})
    ws.freeze_panes = "A4"

    # ---- VOUCHER (cash receipt / payment)
    ws = sheet(wb, "VOUCHER", {"A": 3, "B": 18, "C": 16, "D": 30, "E": 3, "F": 3, "G": 3, "H": 3, "I": 24})
    title(ws, "B1", "سند قبض / سند صرف")
    for r, (lab, nm, v, fmt) in enumerate([
        ("نوع السند", "vcType", MSG["VC_RECEIPT"], None),
        ("التاريخ", "vcDate", datetime.combine(date.today(), datetime.min.time()), "yyyy-mm-dd"),
        ("الخزينة / الحساب النقدي", "vcCash", PAY[1][0], None),
        ("الحساب المقابل (كود)", "vcAcc", "1401", "@"),
        ("الطرف", "vcParty", None, None),
        ("المبلغ", "vcAmount", None, MONEY),
        ("البيان", "vcDesc", None, None),
    ], start=3):
        label(ws, f"B{r}", lab)
        inp(ws, f"C{r}", v, fmt)
        wb_name(wb, nm, f"{q('VOUCHER')}!$C${r}")
    ws.merge_cells("C9:D9")
    inp(ws, "C9:D9")
    ws["D6"].value = '=IF(C6="","",IFERROR(INDEX(tblCOA[اسم الحساب],MATCH(C6,tblCOA[الكود],0)),"؟ كود غير موجود"))'
    ws["D6"].font = Font(name=FONT, bold=True, color=NAVY)
    list_validation(ws, "C3", "lstVcTypes")
    list_validation(ws, "C5", "lstCash")
    list_validation(ws, "C6", 'INDIRECT("tblCOA[الكود]")')
    list_validation(ws, "C7", 'INDIRECT("tblParties[الاسم]")')
    ws["B11"].value = ("سند القبض: من حساب الخزينة إلى الحساب المقابل. سند الصرف: من الحساب المقابل إلى الخزينة. "
                       "مع العملاء (1401) والموردين (2101) لازم تختار الطرف. السند بيروح للاعتماد مباشرة.")
    ws["B11"].font = Font(name=FONT, italic=True, color="595959")
    local_name(ws, "InputArea", f"{q('VOUCHER')}!$C$3:$C$8,{q('VOUCHER')}!$C$9:$D$9")

    # ---- ITEM_CARD
    ws = sheet(wb, "ITEM_CARD", {"A": 3, "B": 11, "C": 11, "D": 16, "E": 20, "F": 28, "G": 11, "H": 11, "I": 11,
                                 "J": 13, "K": 11, "L": 13, "M": 11, "N": 13, "O": 3, "P": 18})
    title(ws, "B1", "كارت صنف (بالمتوسط المرجح)")
    label(ws, "B4", "كود الصنف")
    inp(ws, "C4", all_items[0]["code"] if all_items else None, "@")
    label(ws, "D4", "من")
    inp(ws, "E4", datetime(2026, 4, 1), "yyyy-mm-dd")
    label(ws, "F4", "إلى")
    inp(ws, "G4", datetime(2026, 12, 31), "yyyy-mm-dd")
    ws["C5"].value = '=IF(C4="","",IFERROR(INDEX(tblItems[اسم الصنف],MATCH(C4,tblItems[كود الصنف],0)),"؟"))'
    ws["C5"].font = Font(name=FONT, bold=True, color=NAVY)
    list_validation(ws, "C4", 'INDIRECT("tblItems[كود الصنف]")')
    header_row(ws, 7, 2, ["التاريخ", "رقم الإذن", "النوع", "الطرف / الطلبية", "البيان", "وارد كمية", "صادر كمية",
                          "تكلفة الوحدة", "قيمة الوارد", "رصيد الكمية", "رصيد القيمة", "متوسط التكلفة", "قيمة الصادر"])
    for k, ref in (("icItem", "$C$4"), ("icFrom", "$E$4"), ("icTo", "$G$4")):
        wb_name(wb, k, f"{q('ITEM_CARD')}!{ref}")
    local_name(ws, "ParamArea", ",".join(f"{q('ITEM_CARD')}!{r}" for r in ("$C$4", "$E$4", "$G$4")))
    ws.freeze_panes = "A8"

    # ---- PARTY_BAL
    ws = sheet(wb, "PARTY_BAL", {"A": 3, "B": 30, "C": 12, "D": 16, "E": 16, "F": 16, "G": 9, "H": 3, "I": 3,
                                 "J": 18})
    title(ws, "B1", "أرصدة العملاء والموردين")
    label(ws, "B4", "حتى تاريخ")
    inp(ws, "C4", datetime(2026, 12, 31), "yyyy-mm-dd")
    header_row(ws, 7, 2, ["الاسم", "النوع", "رصيد العملاء (1401)", "رصيد الموردين (2101)", "الرصيد", "طبيعته"])
    wb_name(wb, "pbTo", f"{q('PARTY_BAL')}!$C$4")
    local_name(ws, "ParamArea", f"{q('PARTY_BAL')}!$C$4")
    ws.freeze_panes = "A8"

    # ---- SALES_REP
    ws = sheet(wb, "SALES_REP", {"A": 3, "B": 28, "C": 14, "D": 13, "E": 12, "F": 14, "G": 14, "H": 14, "I": 10,
                                 "J": 14, "K": 3, "L": 3, "M": 18})
    title(ws, "B1", "تقرير المبيعات ومجمل الربح لكل عميل")
    label(ws, "B4", "من")
    inp(ws, "C4", datetime(2026, 4, 1), "yyyy-mm-dd")
    label(ws, "D4", "إلى")
    inp(ws, "E4", datetime(2026, 12, 31), "yyyy-mm-dd")
    header_row(ws, 7, 2, ["العميل", "المبيعات", "المرتجعات", "الخصم", "صافي المبيعات", "تكلفة المبيعات",
                          "مجمل الربح", "نسبة مجمل الربح", "المتحصلات"])
    wb_name(wb, "srFrom", f"{q('SALES_REP')}!$C$4")
    wb_name(wb, "srTo", f"{q('SALES_REP')}!$E$4")
    local_name(ws, "ParamArea", f"{q('SALES_REP')}!$C$4,{q('SALES_REP')}!$E$4")
    ws.freeze_panes = "A8"

    # ---- MOLDS
    mold_code = {m["name"]: f"MD-{i:04d}" for i, m in enumerate(molds, start=1)}
    ws = sheet(wb, "MOLDS", {"A": 3, "B": 10, "C": 30, "D": 22, "E": 12, "F": 9, "G": 11, "H": 13, "I": 13,
                             "J": 10, "K": 13, "L": 12, "M": 14, "N": 14, "O": 13, "P": 24, "Q": 3, "R": 22})
    title(ws, "B1", "سجل الأسطمبات")
    ws["B2"].value = ("التكلفة والعمر المتوقع (بالوحدات) بتتكتب يدوي. الاستهلاك = الوحدات المنتجة × (التكلفة ÷ العمر)، "
                      "للأسطمبات ملك المصنع بس، وبيتقيد شهرياً من شيت تشغيل الأسطمبات.")
    ws["B2"].font = Font(name=FONT, italic=True, color="595959")
    mc = "tblMolds[[#This Row],[الكود]]"
    add_table(ws, "tblMolds", 3, 2, ["الكود", "اسم الأسطمبة", "العميل", "الملكية", "وحدة العد", "تاريخ البدء",
                                     "التكلفة", "العمر المتوقع (وحدات)", "الحالة", "الوحدات المنتجة",
                                     "استهلاك الوحدة", "مجمع الاستهلاك", "صافي القيمة", "المتبقي من العمر", "ملاحظات"],
              [[mold_code[m["name"]], m["name"], m["customer"] or None, MSG["OWN_FACTORY"], m["unit"],
                datetime.combine(m["first"], datetime.min.time()) if m["first"] else None, None, None, "شغالة",
                None, None, None, None, None, "من شيت التكاليف والأوردرات - التكلفة والعمر محتاجين استكمال"]
               for m in molds],
              formulas={10: f"SUMIFS(tblMoldUse[الوحدات],tblMoldUse[كود الأسطمبة],{mc})",
                        11: "IF(AND(N(tblMolds[[#This Row],[التكلفة]])>0,N(tblMolds[[#This Row],[العمر المتوقع (وحدات)]])>0),"
                            "tblMolds[[#This Row],[التكلفة]]/tblMolds[[#This Row],[العمر المتوقع (وحدات)]],0)",
                        12: f'SUMIFS(jeCr,jeCode,"1206",jeNote,{mc}&" - *",jeStatus,stApproved)'
                            f'-SUMIFS(jeDr,jeCode,"1206",jeNote,{mc}&" - *",jeStatus,stApproved)',
                        13: "N(tblMolds[[#This Row],[التكلفة]])-tblMolds[[#This Row],[مجمع الاستهلاك]]",
                        14: 'IF(N(tblMolds[[#This Row],[العمر المتوقع (وحدات)]])=0,"",'
                            'tblMolds[[#This Row],[العمر المتوقع (وحدات)]]-tblMolds[[#This Row],[الوحدات المنتجة]])'},
              text_cols=(1,), formats={6: "yyyy-mm-dd", 7: MONEY, 8: "#,##0", 10: "#,##0", 11: "#,##0.0000",
                                       12: MONEY, 13: MONEY, 14: "#,##0"})
    list_validation(ws, "E4:E1000", "lstOwn")
    list_validation(ws, "F4:F1000", "lstCounter")
    list_validation(ws, "D4:D1000", 'INDIRECT("tblParties[الاسم]")')
    ws.freeze_panes = "A4"

    # ---- MOLD_USE
    ws = sheet(wb, "MOLD_USE", {"A": 3, "B": 11, "C": 10, "D": 30, "E": 12, "F": 22, "G": 11, "H": 9, "I": 11,
                                "J": 11, "K": 30, "L": 3, "M": 24})
    title(ws, "B1", "تشغيل الأسطمبات")
    label(ws, "B2", "شهر القيد")
    inp(ws, "C2", datetime(2026, 7, 1), "yyyy-mm")
    ws["D2"].value = "اختار أي يوم في الشهر واضغط «قيد استهلاك الشهر» (مدين 5209 / دائن 1206)."
    ws["D2"].font = Font(name=FONT, italic=True, color="595959")
    add_table(ws, "tblMoldUse", 4, 2, ["التاريخ", "كود الأسطمبة", "اسم الأسطمبة", "الطلبية", "العميل",
                                       "عدد الضربات", "عدد العيون", "الوحدات", "رقم قيد الاستهلاك", "ملاحظات"],
              [[datetime.combine(dt, datetime.min.time()) if dt else None, mold_code[nm], None, order or None,
                cust or None, qty, cav, None, None, note] for dt, nm, order, cust, qty, cav, note in usage],
              formulas={3: 'IFERROR(INDEX(tblMolds[اسم الأسطمبة],MATCH(tblMoldUse[[#This Row],[كود الأسطمبة]],tblMolds[الكود],0)),"")',
                        8: "N(tblMoldUse[[#This Row],[عدد الضربات]])*MAX(1,N(tblMoldUse[[#This Row],[عدد العيون]]))"},
              text_cols=(2,), formats={1: "yyyy-mm-dd", 6: "#,##0", 7: "0", 8: "#,##0"})
    list_validation(ws, "C5:C5000", 'INDIRECT("tblMolds[الكود]")')
    list_validation(ws, "F5:F5000", 'INDIRECT("tblParties[الاسم]")')
    wb_name(wb, "muMonth", f"{q('MOLD_USE')}!$C$2")
    local_name(ws, "ParamArea", f"{q('MOLD_USE')}!$C$2")
    ws.freeze_panes = "A5"

    # ---- USERS
    ws = sheet(wb, "USERS", {"A": 3, "B": 14, "C": 22, "D": 14, "E": 8, "F": 14, "G": 18, "H": 3, "J": 22})
    title(ws, "B1", "المستخدمين")
    ws["B2"].value = "لإضافة أو تغيير كلمة مرور: اكتبها في «كلمة مرور جديدة» واضغط «حفظ كلمات المرور». كلمة المرور الأولية 1234 لكل المستخدمين، غيّرها فوراً."
    ws["B2"].font = Font(name=FONT, italic=True, color="C00000")
    add_table(ws, "tblUsers", 3, 2, ["المستخدم", "الاسم", "الدور", "نشط", "PassHash", "كلمة مرور جديدة"],
              [[u, n, r, "نعم", pass_hash(u, DEFAULT_PASS), None] for u, n, r in USERS], text_cols=(1,))
    ws.column_dimensions["F"].hidden = True
    list_validation(ws, "D4:D100", 'INDIRECT("tblRoles[الاسم]")')
    list_validation(ws, "E4:E100", "lstYesNo")
    for r in range(4, 101):
        ws[f"G{r}"].number_format = ";;;"

    # ---- PERMS
    ws = sheet(wb, "PERMS", {"A": 3, "B": 11, "C": 22, "D": 10, "E": 10, "F": 10, "G": 12, "H": 3, "J": 22})
    title(ws, "B1", "الصلاحيات")
    ws["B2"].value = "E = يظهر ويتعدل   V = يظهر للعرض بس   N = مخفي   (الماستر شايف كل حاجة دايماً)"
    ws["B2"].font = Font(name=FONT, italic=True, color="595959")
    for col, (_, ar) in zip("DEFG", ROLES):
        ws[f"{col}3"].value = ar
        ws[f"{col}3"].font = Font(name=FONT, bold=True, color=NAVY)
        ws[f"{col}3"].alignment = Alignment(horizontal="center")
    add_table(ws, "tblPerms", 4, 2, ["Key", "الشيت"] + [k for k, _ in ROLES], [list(s) for s in SHEETS])
    list_validation(ws, "E5:G100", '"E,V,N"')
    for r in range(5, 5 + len(SHEETS)):
        ws[f"B{r}"].font = Font(name=FONT, color="808080", size=9)
        for col in "DEFG":
            ws[f"{col}{r}"].alignment = Alignment(horizontal="center")

    # ---- SETTINGS
    ws = sheet(wb, "SETTINGS", {"A": 3, "B": 24, "C": 50, "D": 3, "E": 20})
    title(ws, "B1", "الإعدادات")
    for r, (lab, nm, v, fmt) in enumerate([
        ("اسم الشركة", "setCompany", "NeW StAr", None),
        ("نص الفوتر في الطباعة وPDF", "setFooter", "NeW StAr - مصنع منتجات الرابر", None),
        ("رقم القيد التالي", "setNextJE", n_je + 1, "0"),
        ("رقم إذن المخزن التالي", "setNextST", n_st + 1, "0"),
        ("رقم فاتورة البيع التالية", "setNextSI", n_si + 1, "0"),
        ("بداية الموسم / السنة", "setYearStart", datetime(2026, 4, 1), "yyyy-mm-dd"),
    ], start=3):
        label(ws, f"B{r}", lab)
        inp(ws, f"C{r}", v, fmt)
        wb_name(wb, nm, f"{q('SETTINGS')}!$C${r}")
    local_name(ws, "InputArea", f"{q('SETTINGS')}!$C$3:$C$8")

    # ---- LISTS
    ws = sheet(wb, "LISTS", {"A": 22, "B": 60, "D": 12, "E": 14, "G": 11, "H": 10, "I": 22, "J": 20,
                             "L": 11, "M": 11, "N": 12, "P": 11, "Q": 9, "S": 22, "T": 12, "V": 16, "X": 8})
    add_table(ws, "tblMsg", 1, 1, ["Key", "Text"], [[k, v] for k, v in MSG.items()])
    add_table(ws, "tblRoles", 1, 4, ["Key", "الاسم"], [list(r) for r in ROLES])
    add_table(ws, "tblButtons", 1, 7, ["SheetKey", "Range", "Caption", "Macro"], [list(b) for b in BUTTONS])
    add_table(ws, "tblEditCols", 1, 12, ["SheetKey", "Table", "Columns"], [list(e) for e in EDIT_COLS])
    add_table(ws, "tblPrint", 1, 16, ["SheetKey", "Columns"], [list(p) for p in PRINT_COLS])
    add_table(ws, "tblGroups", 1, 19, ["المجموعة", "حساب المخزون", "بادئة الكود"], [list(g) for g in GROUPS],
              text_cols=(2,))
    ws["V1"].value = "lstApprActions"
    for i, k in enumerate(("ACT_APPROVE", "ACT_REJECT", "ACT_RETURN"), start=2):
        ws[f"V{i}"].value = MSG[k]
    ws["X1"].value = "lstYesNo"
    ws["X2"].value, ws["X3"].value = MSG["YES"], MSG["NO"]
    wb_name(wb, "lstApprActions", f"{q('LISTS')}!$V$2:$V$4")
    wb_name(wb, "lstYesNo", f"{q('LISTS')}!$X$2:$X$3")
    add_table(ws, "tblDocTypes", 1, 26, ["Key", "الاسم", "الحساب المقابل"], [list(t) for t in DOC_TYPES], text_cols=(3,))
    add_table(ws, "tblPay", 1, 30, ["الطريقة", "الحساب"], [list(x) for x in PAY], text_cols=(2,))
    for col, head, vals in (("AG", "lstSaleTypes", [MSG["SL_INVOICE"], MSG["SL_RETURN"]]),
                            ("AH", "lstVcTypes", [MSG["VC_RECEIPT"], MSG["VC_PAYMENT"]]),
                            ("AI", "lstDir", [MSG["DIR_IN"], MSG["DIR_OUT"]]),
                            ("AJ", "lstOwn", [MSG["OWN_FACTORY"], MSG["OWN_CUSTOMER"]]),
                            ("AK", "lstCounter", ["قطعة", "جوز"])):
        ws[f"{col}1"].value = head
        for i, v in enumerate(vals, start=2):
            ws[f"{col}{i}"].value = v
        wb_name(wb, head, f"{q('LISTS')}!${col}$2:${col}${len(vals) + 1}")
    wb_name(wb, "lstDocTypes", f"{q('LISTS')}!$AA$2:$AA$7")
    wb_name(wb, "lstPay", f"{q('LISTS')}!$AD$2:$AD${len(PAY) + 1}")
    wb_name(wb, "lstCash", f"{q('LISTS')}!$AD$3:$AD${len(PAY) + 1}")
    for col in ("Z", "AA", "AB", "AD", "AE", "AG", "AH", "AI", "AJ", "AK"):
        ws.column_dimensions[col].width = 20
    msg_row = list(MSG).index("ST_APPROVED") + 2
    wb_name(wb, "stApproved", f"{q('LISTS')}!$B${msg_row}")

    # tab order: what people use first; WARN active so the .xlsx opens on the instructions
    order = ["WARN", "LOGIN"] + [k for k, *_ in SHEETS if k not in ("WARN", "LOGIN", "LISTS")] + ["LISTS"]
    wb._sheets = [wb[NAME[k]] for k in order]
    wb.active = 0
    wb.save(out_path)
    dr = sum(l[2] for l in lines)
    cr = sum(l[3] for l in lines)
    print(f"items {len(all_items)}, parties {len(parties) + len(hist.new_parties)}, assets {len(assets) + len(new_assets)}, "
          f"opening lines {len(lines)} dr {dr:,.2f} cr {cr:,.2f}; docs {len(docs)}, JE rows {len(je_rows)}, "
          f"stock rows {len(st_rows)}, sales rows {len(sl_rows)}, molds {len(molds)}, usage {len(usage)}")
    if notes_path:
        with open(notes_path, "w", encoding="utf-8") as f:
            f.write("\n".join(f"- {x}" for x in hist.issues) + "\n")
            neg = [(n, q) for n, q in hist.qty.items() if q < -1e-6]
            for n, qty in neg:
                f.write(f"- رصيد سالب في آخر الحركات: {n} = {qty:,.3f}\n")


if __name__ == "__main__":
    build(*sys.argv[1:])
